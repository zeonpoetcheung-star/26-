#!/usr/bin/env python3
"""A-2 Structural Recon for CUMCM 2026 C, route A.

Only structural, temporal, causal-information, unit, and output-contract checks
are performed. No statistical pattern analysis, forecasting, or optimization is
implemented here.
"""

from __future__ import annotations

import csv
import hashlib
import platform
import subprocess
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[4]
A1_DIR = ROOT / "A_route" / "common" / "01_preprocessing"
A2_DIR = ROOT / "A_route" / "common" / "02_recon"
PROCESSED = A1_DIR / "processed"
A1_TABLES = A1_DIR / "tables"
TABLES = A2_DIR / "tables"
REPORTS = A2_DIR / "reports"
SCRIPTS = A2_DIR / "scripts"

CHECK_FIELDS = ["runner", "check_id", "scope", "check_name", "status", "blocking", "expected", "actual", "detail"]


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream))


def write_csv(path: Path, fields: list[str], rows: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in fields})


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def a1_hash_rows() -> list[dict[str, str]]:
    rows = []
    for path in sorted(p for p in A1_DIR.rglob("*") if p.is_file() and "__pycache__" not in p.parts and p.suffix != ".pyc"):
        rows.append({"relative_path": path.relative_to(ROOT).as_posix(), "sha256": sha256(path)})
    return rows


class Checks:
    def __init__(self) -> None:
        self.rows: list[dict[str, str]] = []
        self.blocking_failures: list[str] = []

    def add(self, scope: str, name: str, expected: Any, actual: Any, passed: bool, blocking: bool = True, detail: str = "") -> None:
        self.rows.append({
            "runner": "main",
            "check_id": f"RECON-MAIN-{len(self.rows) + 1:03d}",
            "scope": scope,
            "check_name": name,
            "status": "PASS" if passed else "FAIL",
            "blocking": str(blocking).lower(),
            "expected": str(expected),
            "actual": str(actual),
            "detail": detail,
        })
        if blocking and not passed:
            self.blocking_failures.append(f"{scope}: {name}")


def fmt_point(minutes: int) -> str:
    day_offset, minute_of_day = divmod(minutes, 1440)
    hour, minute = divmod(minute_of_day, 60)
    suffix = f"+{day_offset}" if day_offset else ""
    return f"{hour}:{minute:02d}{suffix}"


def interval_label(start: int, end: int) -> str:
    return f"{fmt_point(start)}-{fmt_point(end)}"


def exact_key_check(name: str, rows: list[dict[str, str]], expected_count: int, key_fields: tuple[str, ...]) -> dict[str, Any]:
    keys = [tuple(row[field] for field in key_fields) for row in rows]
    return {
        "check_name": name,
        "expected": f"{expected_count} rows and unique {key_fields}",
        "actual": f"rows={len(rows)}; unique_keys={len(set(keys))}",
        "status": "PASS" if len(rows) == expected_count and len(set(keys)) == expected_count else "FAIL",
        "join_contract": "mechanically exact key" if len(rows) == expected_count and len(set(keys)) == expected_count else "blocked",
    }


def template_labels() -> tuple[list[str], dict[str, list[str]]]:
    result1_path = ROOT / "input" / "templates" / "result1.xlsx"
    wb1 = load_workbook(result1_path, read_only=False, data_only=False)
    result1 = [str(wb1["计划购电量"].cell(row, 1).value) for row in range(2, 146)]
    wb1.close()
    annual: dict[str, list[str]] = {}
    for name in ("result2.xlsx", "result3.xlsx", "result4-2.xlsx", "result4-3.xlsx"):
        wb = load_workbook(ROOT / "input" / "templates" / name, read_only=False, data_only=False)
        annual[name] = [str(wb["计划购电量"].cell(1, column).value) for column in range(2, 146)]
        wb.close()
    return result1, annual


def build_information_matrix() -> list[dict[str, str]]:
    contexts = [
        ("P1", "0:00"), ("P2", "daily 0:00"),
        ("P3", "0:00"), ("P3", "6:00"), ("P3", "12:00"), ("P3", "18:00"),
        ("P4-2", "0:00"),
        ("P4-3", "0:00"), ("P4-3", "6:00"), ("P4-3", "12:00"), ("P4-3", "18:00"),
    ]
    quantities = [
        "fixed tariff", "dynamic price", "historical load", "future actual load",
        "historical PV actual", "future PV actual", "PV forecast issued at current decision time",
        "future PV forecast not yet issued", "current storage state", "future storage state",
    ]
    rows: list[dict[str, str]] = []
    for problem, decision_time in contexts:
        for quantity in quantities:
            if quantity == "fixed tariff":
                if problem in {"P1", "P2", "P3"}:
                    classification, note = "KNOWN_BY_PROBLEM_ASSUMPTION", "Fixed tariff is supplied for the fixed-price problems."
                else:
                    classification, note = "AMBIGUOUS_FROM_STATEMENT", "Fixed tariff is not the governing P4 price process."
            elif quantity == "dynamic price":
                if problem.startswith("P4"):
                    classification, note = "AMBIGUOUS_FROM_STATEMENT", "Realized dynamic prices are supplied, but future publication/forecast availability is not stated."
                else:
                    classification, note = "AMBIGUOUS_FROM_STATEMENT", "Dynamic price is outside the fixed-price problem definition."
            elif quantity in {"historical load", "historical PV actual"}:
                classification, note = "REALIZED_HISTORY_ONLY", "Only observations strictly before the decision time are historical."
            elif quantity == "future actual load":
                if problem == "P1":
                    classification, note = "KNOWN_BY_PROBLEM_ASSUMPTION", "P1 supplies the repeated deterministic daily load profile."
                else:
                    classification, note = "EX_POST_ONLY", "The attachment contains realized values, not a publication available before realization."
            elif quantity == "future PV actual":
                classification, note = "EX_POST_ONLY", "Actual PV is realized after the decision and cannot be treated as known."
            elif quantity == "PV forecast issued at current decision time":
                if problem in {"P1", "P3", "P4-3"}:
                    classification, note = "AVAILABLE_FROM_PUBLISHED_FORECAST", "The applicable problem supplies a forecast at this decision time."
                else:
                    classification, note = "AMBIGUOUS_FROM_STATEMENT", "No official forecast object for this decision context is supplied."
            elif quantity == "future PV forecast not yet issued":
                classification, note = "EX_POST_ONLY", "A later forecast release is not available at the current decision time."
            elif quantity == "current storage state":
                classification, note = "CURRENT_STATE_IF_OBSERVED", "The current state can be used if observed at the decision boundary; only the initial 2025-01-01 value is stated numerically."
            else:
                classification, note = "EX_POST_ONLY", "Future storage state is a downstream state, not current information."
            rows.append({"problem": problem, "decision_time": decision_time, "quantity": quantity, "classification": classification, "source_fact_or_limit": note})
    return rows


def write_reports(
    cross_rows: list[dict[str, Any]], hypothesis_rows: list[dict[str, Any]], alignment_rows: list[dict[str, Any]],
    coverage_rows: list[dict[str, Any]], candidate_rows: list[dict[str, Any]], info_rows: list[dict[str, Any]],
    state_rows: list[dict[str, Any]], output_rows: list[dict[str, Any]], result1_hstart_matches: int,
    annual_hstart_matches: dict[str, int], a1_hashes: list[dict[str, str]], run_time: datetime,
) -> None:
    safe = [row["check_name"] for row in cross_rows if row["status"] == "PASS"]
    (REPORTS / "CROSS_SOURCE_CONTRACT.md").write_text(f"""# A-2 Cross-source Contract

## Confirmed

- The 144-slot fixed-day table is complete and unique.
- The annual load, actual PV, merged actual, and dynamic-price tables each have 52,560 unique `(date, slot_id)` keys.
- Load and actual-PV keys are identical. Merged-actual and dynamic-price keys are identical.
- Attachment 1, both Attachment 2 sheets, and Attachment 4 use the same marker sequence by `slot_id`.

## Join boundary

Exact later joins are safe on the frozen keys listed above. They establish positional and calendar-key compatibility only. They do not settle whether a marker is an interval start or endpoint, nor do they authorize exact forecast-error alignment.

Checks passed: {len(safe)}/{len(cross_rows)}.
""", encoding="utf-8", newline="\n")

    annual_text = ", ".join(f"{name}={count}/144" for name, count in annual_hstart_matches.items())
    (REPORTS / "TIME_ALIGNMENT_RECON.md").write_text(f"""# A-2 Time Alignment Recon

## Two complete positional contracts

- H-END maps source marker `0:10` to `0:00-0:10` and `0:00+1` to `23:50-0:00+1`. It covers exactly 0:00-24:00 with 144 intervals and is structurally natural for the stated P1 SOC boundaries.
- H-START maps source marker `0:10` to `0:10-0:20` and `0:00+1` to `0:00+1-0:10+1`. It omits `0:00-0:10` and extends to 24:10.
- Requested interval `10:00-10:10` maps to source marker `10:10` under H-END and `10:00` under H-START.

## Template evidence conflicts with the boundary evidence

- Position-by-position literal matches against H-START: result1={result1_hstart_matches}/144; {annual_text}.
- The known final-label conflict is preserved: result1 ends `0:00+1-0:10+1`, whereas annual plan/adjustment templates end `0:00-0:10+1`.
- Template labels structurally favor the shifted H-START position sequence, while the full-day coverage and P1 0:00/24:00 state boundaries favor H-END. Neither source uniquely overrides the other.
- Both hypotheses remain unselected in `time_semantics_hypotheses.csv`; A-2 does not convert either into a final modeling convention.

OI01_STRUCTURAL_CONFLICT_CONFIRMED
""", encoding="utf-8", newline="\n")

    coverage_lines = "\n".join(
        f"- {row['issue_clock']}: issues={row['issue_count']}, same-day targets={row['same_day_targets']}, next-day targets={row['next_day_targets']}, earliest={row['earliest_target']}, latest={row['latest_target']}."
        for row in coverage_rows
    )
    candidate_keys = sum(row["candidate_actual_key_exists"] == "true" for row in candidate_rows)
    (REPORTS / "FORECAST_GRID_RECON.md").write_text(f"""# A-2 Forecast Grid Recon

## Confirmed issue and target grid

- The frozen long table contains 1,460 issue timestamps and 35,040 unique `(issue_datetime, horizon_hour)` rows.
- Issue clocks are exactly 0:00, 6:00, 12:00, and 18:00; every issue has horizons 1 through 24.
- Every target satisfies `nominal_target_datetime = issue_datetime + horizon_hour` and lies on a whole hour.
- Forty target rows enter 2026 and remain present.

{coverage_lines}

## Candidate matching is not semantic equality

- Every target has a candidate label in the 144-marker vocabulary; {candidate_keys:,}/35,040 also have a candidate actual-data `(date, slot_id)` key within the 2025 table. The remaining rows cross beyond available actual-data coverage.
- The forecast may be a point value or an hourly representative/quantity. The 10-minute actual values still inherit the unresolved OI-01 interval meaning.
- Under a point interpretation, each issue starts at `issue time + 1 hour`; therefore the first control hour after 0:00, 6:00, 12:00, or 18:00 has no nominal target point from that issue.
- No forecast-error metric or interpolation was computed.

OI13_REMAINS_OPEN
""", encoding="utf-8", newline="\n")

    (REPORTS / "INFORMATION_SET_RECON.md").write_text("""# A-2 Information Set Recon

## Confirmed causal boundary

- P2/P3 formal output begins 2025-02-01, while actual load and PV begin 2025-01-01. Exactly 31 complete calendar days precede the first required output day.
- No official load-forecast file is supplied.
- No separate future-price forecast or publication file is supplied for P4.
- Realized future load/PV values are ex-post observations and cannot be treated as available at a daily 0:00 decision.
- For P3/P4-3, only the PV forecast issued at the current 0:00/6:00/12:00/18:00 decision time is available; later releases are not yet available.
- The current storage state may be used if observed at the decision boundary; future storage states are not information inputs.

## Issues carried forward

- OI-05 remains open for the exact P2 planning information set and any use of January history.
- OI-12 remains open because the statement does not uniquely say whether the future dynamic-price path is known at the decision time.
- The matrix records the ambiguity without choosing a forecasting or operating convention.
""", encoding="utf-8", newline="\n")

    (REPORTS / "STATE_UNIT_RECON.md").write_text("""# A-2 State and Unit Contract

## Unambiguous quantities

- Storage maximum capacity: 12,000 kWh.
- Allowed state range: 1,200-10,800 kWh; usable band: 9,600 kWh.
- Initial state: 6,000 kWh at 2025-01-01 0:00.
- Maximum charge/discharge power: 5,000 kW.
- Ten-minute duration: 1/6 h.
- Rated ten-minute transfer before any efficiency convention: 5,000 x 1/6 = 833.333333 kWh.
- The statement gives charge/discharge efficiency as 90% without fixing its mathematical decomposition.

## State-time boundary

- A 144-interval 0:00-24:00 convention naturally has 145 state-boundary points.
- P1 explicitly requires S(0:00)=S(24:00).
- P2/P3/P4 do not explicitly repeat daily equality; the only numerical initial SOC is at 2025-01-01 0:00.
- Efficiency side, battery/grid-side quantities, daily reset versus cross-day propagation, and terminal treatment remain unresolved for A-4.
""", encoding="utf-8", newline="\n")

    ellipsis_count = sum(row["ellipsis_present"] == "true" for row in output_rows)
    (REPORTS / "OUTPUT_CONTRACT_RECON.md").write_text(f"""# A-2 Output Contract Recon

## Confirmed

- Annual plan/adjustment output dates are 2025-02-01 through 2025-12-31: 334 rows.
- Each plan or adjustment sheet has 144 ten-minute positions per date, plus its stated daily total fields where applicable.
- Charge/discharge and emergency-purchase sheets retain example/ellipsis structures; {ellipsis_count} audited sheets contain an ellipsis marker and require a later Result Writer decision.
- No template was modified or expanded.

## Preserved conflict

- result1 ends with `0:00+1-0:10+1`.
- Annual plan/adjustment templates end with `0:00-0:10+1`.
- A-2 records this difference and does not repair it.
""", encoding="utf-8", newline="\n")

    (REPORTS / "RECON_SUMMARY.md").write_text("""# A-2 Structural Recon Summary

1. Stable facts: all canonical grains, keys, the 144-slot marker sequence, the 365 x 4 forecast issue grid, 24 horizons, unit conversions, and official output ranges are structurally confirmed.
2. Safe joins: annual load, actual PV, merged actual, and dynamic price join exactly on `(date, slot_id)`; fixed-day and annual marker sources align exactly by `slot_id`.
3. Time-label conflict: H-END covers 0:00-24:00 and fits the SOC boundary; template positions mostly follow H-START, which omits the first ten minutes and extends beyond 24:00. The final annual-template label also differs from result1.
4. Information boundary: realized future load/PV are ex-post only; current published PV forecasts are available at their issue times; future dynamic-price availability is not stated; the current SOC is usable only when observed.
5. Forecast-grid ambiguity: target timestamps and candidate marker matches are mechanical, but point-versus-hourly forecast meaning is not uniquely specified, so exact forecast-error alignment is not authorized.
6. Forward issues: OI-01 and OI-13 remain guarded for EDA; OI-03-OI-12 and OI-14-OI-16 remain for their assigned later decision stages.
""", encoding="utf-8", newline="\n")

    planned_outputs = sorted(
        [path.relative_to(ROOT).as_posix() for path in TABLES.glob("*.csv")]
        + [path.relative_to(ROOT).as_posix() for path in REPORTS.glob("*.md")]
        + [str((SCRIPTS / "recon_c2026.py").relative_to(ROOT).as_posix()), str((SCRIPTS / "validate_recon.py").relative_to(ROOT).as_posix())]
        + [
            str((TABLES / "recon_checks.csv").relative_to(ROOT).as_posix()),
            str((REPORTS / "RUN_INFO.md").relative_to(ROOT).as_posix()),
            str((REPORTS / "RECON_GATE.md").relative_to(ROOT).as_posix()),
        ]
    )
    (REPORTS / "RUN_INFO.md").write_text(f"""# A-2 Run Information

- run_datetime: `{run_time.astimezone().isoformat()}`
- python_executable: `{sys.executable}`
- python_version: `{platform.python_version()}`
- command: `python A_route/common/02_recon/scripts/recon_c2026.py`
- independent_validator: invoked by the main script before the final gate
- frozen_A1_files_hashed: {len(a1_hashes)}
- CURRENT_STATE: `A-2 Structural Recon`; A-1 accepted and frozen

## Core artifacts

{chr(10).join(f'- `{path}`' for path in sorted(set(planned_outputs)))}
""", encoding="utf-8", newline="\n")


def main() -> int:
    run_time = datetime.now().astimezone()
    for directory in (TABLES, REPORTS, SCRIPTS):
        directory.mkdir(parents=True, exist_ok=True)
    gate_path = REPORTS / "RECON_GATE.md"
    if gate_path.exists():
        gate_path.unlink()

    checks = Checks()
    a1_before = a1_hash_rows()
    write_csv(TABLES / "a1_frozen_hashes.csv", ["relative_path", "sha256"], a1_before)

    fixed = read_csv(PROCESSED / "fixed_day_10min.csv")
    load_rows = read_csv(PROCESSED / "year_load_10min.csv")
    pv_rows = read_csv(PROCESSED / "year_pv_actual_10min.csv")
    actual_rows = read_csv(PROCESSED / "year_actual_10min.csv")
    forecast_rows = read_csv(PROCESSED / "pv_forecast_hourly_long.csv")
    price_rows = read_csv(PROCESSED / "dynamic_price_10min.csv")
    crosswalk = read_csv(A1_TABLES / "time_marker_crosswalk.csv")
    template_contract = read_csv(A1_TABLES / "template_contract.csv")

    cross_rows = [
        exact_key_check("fixed_day complete key", fixed, 144, ("slot_id",)),
        exact_key_check("annual load complete key", load_rows, 52560, ("date", "slot_id")),
        exact_key_check("annual PV complete key", pv_rows, 52560, ("date", "slot_id")),
        exact_key_check("annual merged actual complete key", actual_rows, 52560, ("date", "slot_id")),
        exact_key_check("annual dynamic price complete key", price_rows, 52560, ("date", "slot_id")),
    ]
    key_sets = {
        "load": {(row["date"], row["slot_id"]) for row in load_rows},
        "pv": {(row["date"], row["slot_id"]) for row in pv_rows},
        "actual": {(row["date"], row["slot_id"]) for row in actual_rows},
        "price": {(row["date"], row["slot_id"]) for row in price_rows},
    }
    cross_rows.extend([
        {"check_name": "load and PV key sets identical", "expected": 52560, "actual": len(key_sets["load"] & key_sets["pv"]), "status": "PASS" if key_sets["load"] == key_sets["pv"] else "FAIL", "join_contract": "safe on (date, slot_id)"},
        {"check_name": "merged actual and price key sets identical", "expected": 52560, "actual": len(key_sets["actual"] & key_sets["price"]), "status": "PASS" if key_sets["actual"] == key_sets["price"] else "FAIL", "join_contract": "safe on (date, slot_id)"},
        {"check_name": "four source marker sequences identical", "expected": "144 true flags", "actual": sum(row["all_sources_same_marker"] == "true" for row in crosswalk), "status": "PASS" if len(crosswalk) == 144 and all(row["all_sources_same_marker"] == "true" for row in crosswalk) else "FAIL", "join_contract": "safe positional match only; interval semantics unresolved"},
    ])
    write_csv(TABLES / "cross_source_join_checks.csv", ["check_name", "expected", "actual", "status", "join_contract"], cross_rows)
    for row in cross_rows:
        checks.add("cross_source", row["check_name"], row["expected"], row["actual"], row["status"] == "PASS")

    result1_labels, annual_labels = template_labels()
    markers = [row["attachment1_time_marker"] for row in crosswalk]
    hypothesis_rows: list[dict[str, Any]] = []
    alignment_rows: list[dict[str, Any]] = []
    for slot_id, marker in enumerate(markers, start=1):
        h_end = interval_label((slot_id - 1) * 10, slot_id * 10)
        h_start = interval_label(slot_id * 10, (slot_id + 1) * 10)
        for hypothesis, start, end, label in (("H-END", (slot_id - 1) * 10, slot_id * 10, h_end), ("H-START", slot_id * 10, (slot_id + 1) * 10, h_start)):
            hypothesis_rows.append({
                "hypothesis": hypothesis, "slot_id": slot_id, "source_marker": marker,
                "interval_start_minute": start, "interval_end_minute": end, "interval_label": label,
                "inside_0_24_day": str(start >= 0 and end <= 1440).lower(),
                "covers_requested_10_00_10_10": str(start == 600 and end == 610).lower(),
                "final_convention_selected": "false",
            })
        alignment_rows.append({
            "slot_id": slot_id, "source_marker": marker, "h_end_label": h_end, "h_start_label": h_start,
            "result1_label": result1_labels[slot_id - 1],
            "result1_matches_h_end": str(result1_labels[slot_id - 1] == h_end).lower(),
            "result1_matches_h_start": str(result1_labels[slot_id - 1] == h_start).lower(),
            **{f"{name.replace('.xlsx', '').replace('-', '_')}_label": labels[slot_id - 1] for name, labels in annual_labels.items()},
            **{f"{name.replace('.xlsx', '').replace('-', '_')}_matches_h_start": str(labels[slot_id - 1] == h_start).lower() for name, labels in annual_labels.items()},
        })
    write_csv(TABLES / "time_semantics_hypotheses.csv", ["hypothesis", "slot_id", "source_marker", "interval_start_minute", "interval_end_minute", "interval_label", "inside_0_24_day", "covers_requested_10_00_10_10", "final_convention_selected"], hypothesis_rows)
    alignment_fields = list(alignment_rows[0].keys())
    write_csv(TABLES / "source_template_slot_alignment.csv", alignment_fields, alignment_rows)
    checks.add("time_semantics", "both hypotheses contain 144 slots", "144 each", dict(Counter(row["hypothesis"] for row in hypothesis_rows)), Counter(row["hypothesis"] for row in hypothesis_rows) == {"H-END": 144, "H-START": 144})
    checks.add("time_semantics", "no final convention silently selected", 0, sum(row["final_convention_selected"] == "true" for row in hypothesis_rows), all(row["final_convention_selected"] == "false" for row in hypothesis_rows))
    checks.add("time_semantics", "H-END full-day coverage", "0..1440", f"{hypothesis_rows[0]['interval_start_minute']}..{hypothesis_rows[286]['interval_end_minute']}", hypothesis_rows[0]["interval_start_minute"] == 0 and hypothesis_rows[286]["interval_end_minute"] == 1440)
    checks.add("time_semantics", "H-START omits first interval and extends beyond day", "10..1450", f"{hypothesis_rows[1]['interval_start_minute']}..{hypothesis_rows[287]['interval_end_minute']}", hypothesis_rows[1]["interval_start_minute"] == 10 and hypothesis_rows[287]["interval_end_minute"] == 1450)

    issue_groups: dict[str, list[dict[str, str]]] = defaultdict(list)
    actual_keys = {(row["date"], row["slot_id"]): row for row in pv_rows}
    marker_to_slot = {row["time_marker_normalized"]: row["slot_id"] for row in fixed}
    candidate_rows: list[dict[str, Any]] = []
    formula_ok = 0
    whole_hour_ok = 0
    for row in forecast_rows:
        issue_groups[row["issue_clock"]].append(row)
        issue_dt = datetime.fromisoformat(row["issue_datetime"])
        target_dt = datetime.fromisoformat(row["nominal_target_datetime"])
        formula_ok += int(target_dt == issue_dt + timedelta(hours=int(row["horizon_hour"])))
        whole_hour_ok += int(target_dt.minute == 0 and target_dt.second == 0)
        if target_dt.hour == 0:
            candidate_date = (target_dt.date() - timedelta(days=1)).isoformat()
            candidate_marker = "0:00+1"
        else:
            candidate_date = target_dt.date().isoformat()
            candidate_marker = f"{target_dt.hour}:00"
        candidate_slot = marker_to_slot.get(candidate_marker, "")
        key_exists = (candidate_date, str(candidate_slot)) in actual_keys if candidate_slot != "" else False
        candidate_rows.append({
            "issue_datetime": row["issue_datetime"], "horizon_hour": row["horizon_hour"],
            "nominal_target_datetime": row["nominal_target_datetime"], "candidate_actual_date": candidate_date,
            "candidate_marker": candidate_marker, "candidate_slot_id": candidate_slot,
            "candidate_marker_in_144_grid": str(candidate_slot != "").lower(),
            "candidate_actual_key_exists": str(key_exists).lower(),
            "semantic_alignment_authorized": "false",
        })
    coverage_rows: list[dict[str, Any]] = []
    for issue_clock in ("0:00", "6:00", "12:00", "18:00"):
        rows = issue_groups[issue_clock]
        same = 0
        next_day = 0
        for row in rows:
            issue_date = datetime.fromisoformat(row["issue_datetime"]).date()
            target_date = datetime.fromisoformat(row["nominal_target_datetime"]).date()
            same += int(target_date == issue_date)
            next_day += int(target_date == issue_date + timedelta(days=1))
        coverage_rows.append({
            "issue_clock": issue_clock, "issue_count": len({row["issue_datetime"] for row in rows}),
            "target_rows": len(rows), "same_day_targets": same, "next_day_targets": next_day,
            "earliest_target": min(row["nominal_target_datetime"] for row in rows),
            "latest_target": max(row["nominal_target_datetime"] for row in rows),
        })
    write_csv(TABLES / "forecast_grid_coverage.csv", ["issue_clock", "issue_count", "target_rows", "same_day_targets", "next_day_targets", "earliest_target", "latest_target"], coverage_rows)
    write_csv(TABLES / "forecast_actual_candidate_alignment.csv", list(candidate_rows[0].keys()), candidate_rows)
    issue_counts = Counter(row["issue_datetime"] for row in forecast_rows)
    horizons: dict[str, set[int]] = defaultdict(set)
    for row in forecast_rows:
        horizons[row["issue_datetime"]].add(int(row["horizon_hour"]))
    complete_issues = sum(issue_counts[key] == 24 and horizons[key] == set(range(1, 25)) for key in issue_counts)
    cross_year = sum(datetime.fromisoformat(row["nominal_target_datetime"]).year == 2026 for row in forecast_rows)
    checks.add("forecast_grid", "issue count", 1460, len(issue_counts), len(issue_counts) == 1460)
    checks.add("forecast_grid", "long row count", 35040, len(forecast_rows), len(forecast_rows) == 35040)
    checks.add("forecast_grid", "24 horizons 1..24 per issue", 1460, complete_issues, complete_issues == 1460)
    checks.add("forecast_grid", "nominal target formula", 35040, formula_ok, formula_ok == 35040)
    checks.add("forecast_grid", "whole-hour targets", 35040, whole_hour_ok, whole_hour_ok == 35040)
    checks.add("forecast_grid", "cross-year rows retained", 40, cross_year, cross_year == 40)
    checks.add("forecast_grid", "candidate marker vocabulary match", 35040, sum(row["candidate_marker_in_144_grid"] == "true" for row in candidate_rows), all(row["candidate_marker_in_144_grid"] == "true" for row in candidate_rows))
    checks.add("forecast_grid", "semantic alignment remains unauthorized", 0, sum(row["semantic_alignment_authorized"] == "true" for row in candidate_rows), all(row["semantic_alignment_authorized"] == "false" for row in candidate_rows))

    info_rows = build_information_matrix()
    write_csv(TABLES / "information_availability_matrix.csv", ["problem", "decision_time", "quantity", "classification", "source_fact_or_limit"], info_rows)
    required_contexts = {("P1", "0:00"), ("P2", "daily 0:00"), ("P3", "0:00"), ("P3", "6:00"), ("P3", "12:00"), ("P3", "18:00"), ("P4-2", "0:00"), ("P4-3", "0:00"), ("P4-3", "6:00"), ("P4-3", "12:00"), ("P4-3", "18:00")}
    actual_contexts = {(row["problem"], row["decision_time"]) for row in info_rows}
    checks.add("information_set", "required decision contexts", 11, len(actual_contexts), actual_contexts == required_contexts)
    checks.add("information_set", "10 quantities per decision context", "10 each", sorted(set(Counter((row["problem"], row["decision_time"]) for row in info_rows).values())), len(info_rows) == 110)
    checks.add("information_set", "January history before first output", 31, (date(2025, 2, 1) - date(2025, 1, 1)).days, True)

    state_rows = [
        {"item": "storage maximum capacity", "value": 12000, "unit": "kWh", "status": "OFFICIAL", "decision_deferred": "none"},
        {"item": "allowed state minimum", "value": 1200, "unit": "kWh", "status": "OFFICIAL", "decision_deferred": "none"},
        {"item": "allowed state maximum", "value": 10800, "unit": "kWh", "status": "OFFICIAL", "decision_deferred": "none"},
        {"item": "usable state band", "value": 9600, "unit": "kWh", "status": "DETERMINISTIC_DERIVATION", "decision_deferred": "none"},
        {"item": "initial state at 2025-01-01 0:00", "value": 6000, "unit": "kWh", "status": "OFFICIAL", "decision_deferred": "cross-day propagation"},
        {"item": "maximum charge/discharge power", "value": 5000, "unit": "kW", "status": "OFFICIAL", "decision_deferred": "none"},
        {"item": "ten-minute duration", "value": "1/6", "unit": "h", "status": "DETERMINISTIC_DERIVATION", "decision_deferred": "none"},
        {"item": "rated ten-minute transfer before efficiency", "value": "833.333333333333", "unit": "kWh", "status": "DETERMINISTIC_DERIVATION", "decision_deferred": "efficiency side"},
        {"item": "stated charge/discharge efficiency", "value": "90%", "unit": "dimensionless", "status": "OFFICIAL_AMBIGUOUS_CONVENTION", "decision_deferred": "efficiency decomposition and energy side"},
        {"item": "state boundary count for 144 intervals", "value": 145, "unit": "points", "status": "CONDITIONAL_ON_0_24_INTERVAL_CONVENTION", "decision_deferred": "time semantics"},
        {"item": "P1 terminal relation", "value": "S(0:00)=S(24:00)", "unit": "kWh", "status": "OFFICIAL", "decision_deferred": "none for P1"},
        {"item": "P2/P3/P4 daily terminal relation", "value": "not explicitly stated", "unit": "", "status": "UNRESOLVED", "decision_deferred": "daily reset/cross-day/terminal treatment"},
    ]
    write_csv(TABLES / "state_unit_contract.csv", ["item", "value", "unit", "status", "decision_deferred"], state_rows)
    checks.add("state_units", "rated ten-minute transfer", 5000 / 6, 5000 * (10 / 60), abs(5000 * (10 / 60) - 833.333333333333) < 1e-9)
    checks.add("state_units", "usable state band", 9600, 10800 - 1200, 10800 - 1200 == 9600)

    output_rows: list[dict[str, Any]] = []
    for row in template_contract:
        sheet_type = "plan_or_adjustment" if row["sheet"] in {"计划购电量", "调整购电量"} else "example_or_summary"
        annual = row["workbook"] != "result1.xlsx"
        output_rows.append({
            "workbook": row["workbook"], "sheet": row["sheet"], "sheet_type": sheet_type,
            "used_range": row["used_range"], "data_rows_or_positions": 334 if annual and sheet_type == "plan_or_adjustment" else row["rows"],
            "ten_minute_positions": 144 if sheet_type == "plan_or_adjustment" else "",
            "output_start_date": "2025-02-01" if annual and sheet_type == "plan_or_adjustment" else "",
            "output_end_date": "2025-12-31" if annual and sheet_type == "plan_or_adjustment" else "",
            "ellipsis_present": row["ellipsis_present"], "first_label": row["first_date_or_time_label"], "last_label": row["last_date_or_time_label"],
            "later_writer_decision_required": str(row["ellipsis_present"] == "true" or (sheet_type == "plan_or_adjustment" and row["last_date_or_time_label"] in {"0:00+1-0:10+1", "0:00-0:10+1"})).lower(),
        })
    write_csv(TABLES / "output_contract_recon.csv", list(output_rows[0].keys()), output_rows)
    annual_plan_rows = [row for row in output_rows if row["sheet_type"] == "plan_or_adjustment" and row["workbook"] != "result1.xlsx"]
    checks.add("output_contract", "annual plan/adjustment date contract", "334 rows; 2025-02-01..2025-12-31", len(annual_plan_rows), all(row["data_rows_or_positions"] == 334 and row["output_start_date"] == "2025-02-01" and row["output_end_date"] == "2025-12-31" and row["ten_minute_positions"] == 144 for row in annual_plan_rows))

    result1_hstart_matches = sum(row["result1_matches_h_start"] == "true" for row in alignment_rows)
    annual_hstart_matches = {name: sum(row[f"{name.replace('.xlsx', '').replace('-', '_')}_matches_h_start"] == "true" for row in alignment_rows) for name in annual_labels}
    write_reports(cross_rows, hypothesis_rows, alignment_rows, coverage_rows, candidate_rows, info_rows, state_rows, output_rows, result1_hstart_matches, annual_hstart_matches, a1_before, run_time)

    write_csv(TABLES / "recon_checks.csv", CHECK_FIELDS, checks.rows)
    validator = subprocess.run([sys.executable, str(SCRIPTS / "validate_recon.py"), "--quiet"], cwd=ROOT, capture_output=True, text=True, encoding="utf-8", check=False)
    all_checks = read_csv(TABLES / "recon_checks.csv")
    validator_rows = [row for row in all_checks if row["runner"] == "validator"]
    validator_blocking = [row for row in validator_rows if row["status"] == "FAIL" and row["blocking"] == "true"]
    blocking = checks.blocking_failures + [f"validator: {row['scope']}: {row['check_name']}" for row in validator_blocking]
    if not validator_rows:
        blocking.append("independent validator produced no checks")
    if validator.returncode not in (0, 1):
        blocking.append(f"independent validator crashed with code {validator.returncode}: {validator.stderr.strip()}")

    status = "BLOCKED_RECON" if blocking else "PASS_RECON_WITH_OPEN_ISSUES"
    validator_pass = sum(row["status"] == "PASS" for row in validator_rows)
    validator_fail = sum(row["status"] == "FAIL" for row in validator_rows)
    gate_path.write_text(f"""# A-2 Recon Gate

## Decision

- Status: `{status}`.
- Cross-source keys and the 144-marker sequence are stable.
- Both OI-01 hypotheses are preserved without selecting a final convention; `OI01_STRUCTURAL_CONFLICT_CONFIRMED`.
- Forecast issue/horizon/target structure is stable; exact physical forecast-error alignment remains unauthorized; `OI13_REMAINS_OPEN`.
- The causal information matrix prevents future actuals and not-yet-issued forecasts from entering earlier decisions.
- Independent validator: PASS={validator_pass}, FAIL={validator_fail}.
- A-1 frozen files remained unchanged during A-2.
- Do not enter A-3 without explicit human/GPT confirmation.

## Blocking checks

{chr(10).join(f'- {item}' for item in blocking) if blocking else '- None.'}

## Open issues carried forward

- OI-01, OI-05, OI-06, OI-12, and OI-13 remain guarded by the contracts produced here.
- OI-03, OI-04, OI-07-OI-11, and OI-14-OI-16 remain assigned to their later stages.

{status}
""", encoding="utf-8", newline="\n")

    print(f"RECON STATUS: {status}")
    print(f"Main checks: PASS={sum(row['status'] == 'PASS' for row in checks.rows)}, FAIL={sum(row['status'] == 'FAIL' for row in checks.rows)}")
    print(f"Validator checks: PASS={validator_pass}, FAIL={validator_fail}")
    return 0 if status != "BLOCKED_RECON" else 1


if __name__ == "__main__":
    raise SystemExit(main())

#!/usr/bin/env python3
"""Independent validator for CUMCM 2026 C, route A, A-1 outputs.

This script reopens original files and generated CSVs. It does not import the
preprocessing implementation and it deliberately runs before the final gate is
written.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
from collections import Counter, defaultdict
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[4]
STAGE_DIR = ROOT / "A_route" / "common" / "01_preprocessing"
PROCESSED_DIR = STAGE_DIR / "processed"
TABLES_DIR = STAGE_DIR / "tables"
REPORTS_DIR = STAGE_DIR / "reports"

CHECK_FIELDS = [
    "runner", "check_id", "scope", "check_name", "status", "blocking",
    "expected", "actual", "detail",
]


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream))


def normalize_time(value: Any) -> str:
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return f"{value.hour}:{value.minute:02d}"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        minutes = int(round((float(value) % 1) * 1440)) % 1440
        return f"{minutes // 60}:{minutes % 60:02d}"
    return str(value).strip() if value is not None else ""


def parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    return None


def expected_dates() -> list[date]:
    start = date(2025, 1, 1)
    return [start + timedelta(days=offset) for offset in range(365)]


def trailing_cells_ooxml(path: Path) -> dict[str, tuple[str | None, str, str]]:
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ns = {"m": main_ns}
    with ZipFile(path) as archive:
        shared_root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
        shared = ["".join(t.text or "" for t in si.iterfind(".//m:t", ns)) for si in shared_root.findall("m:si", ns)]
        sheet_root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        result: dict[str, tuple[str | None, str, str]] = {}
        for block in range(365):
            start_row = 2 + block * 4
            for row in range(start_row + 1, start_row + 4):
                coord = f"A{row}"
                cell = sheet_root.find(f".//m:c[@r='{coord}']", ns)
                if cell is None:
                    result[coord] = (None, "", "")
                    continue
                cell_type = cell.get("t")
                raw_v = cell.findtext("m:v", default="", namespaces=ns)
                resolved = shared[int(raw_v)] if cell_type == "s" and raw_v.isdigit() else raw_v
                result[coord] = (cell_type, raw_v, resolved)
        return result


class Validator:
    def __init__(self) -> None:
        self.rows: list[dict[str, str]] = []
        self.blocking_failures: list[str] = []

    def check(
        self, scope: str, name: str, expected: Any, actual: Any,
        passed: bool, blocking: bool = True, detail: str = "",
    ) -> None:
        self.rows.append(
            {
                "runner": "validator",
                "check_id": f"VALIDATOR-{len(self.rows) + 1:03d}",
                "scope": scope,
                "check_name": name,
                "status": "PASS" if passed else "FAIL",
                "blocking": str(blocking).lower(),
                "expected": str(expected),
                "actual": str(actual),
                "detail": detail,
            }
        )
        if not passed and blocking:
            self.blocking_failures.append(f"{scope}: {name}")


def validate_csv_keyed(
    path: Path, expected_count: int, key_fields: tuple[str, ...],
    validator: Validator, scope: str,
) -> list[dict[str, str]]:
    rows = read_rows(path)
    keys = [tuple(row[field] for field in key_fields) for row in rows]
    validator.check(scope, "CSV row count", expected_count, len(rows), len(rows) == expected_count)
    validator.check(scope, f"unique key {key_fields}", expected_count, len(set(keys)), len(set(keys)) == expected_count)
    return rows


def rewrite_validation_checks(validator_rows: list[dict[str, str]]) -> None:
    path = TABLES_DIR / "validation_checks.csv"
    existing = read_rows(path) if path.exists() else []
    main_rows = [row for row in existing if row.get("runner") != "validator"]
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=CHECK_FIELDS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(main_rows + validator_rows)


def validate(quiet: bool = False) -> int:
    validator = Validator()
    gate_path = REPORTS_DIR / "PREPROCESSING_GATE.md"
    validator.check("pipeline", "gate absent before independent validator", False, gate_path.exists(), not gate_path.exists())

    manifest = read_rows(TABLES_DIR / "data_manifest.csv")
    validator.check("manifest", "10 official input files", 10, len(manifest), len(manifest) == 10)
    for row in manifest:
        current_hash = file_hash(ROOT / row["relative_path"])
        validator.check("manifest", f"SHA-256 {row['file_name']}", row["sha256"], current_hash, current_hash.lower() == row["sha256"].lower())

    fixed = validate_csv_keyed(PROCESSED_DIR / "fixed_day_10min.csv", 144, ("slot_id",), validator, "fixed_day_10min")
    fixed_slots = [int(row["slot_id"]) for row in fixed]
    validator.check("fixed_day_10min", "slot_id exactly 1..144", "1..144", f"{min(fixed_slots)}..{max(fixed_slots)}", fixed_slots == list(range(1, 145)))

    load_rows = validate_csv_keyed(PROCESSED_DIR / "year_load_10min.csv", 52560, ("date", "slot_id"), validator, "year_load_10min")
    pv_rows = validate_csv_keyed(PROCESSED_DIR / "year_pv_actual_10min.csv", 52560, ("date", "slot_id"), validator, "year_pv_actual_10min")
    merged_rows = validate_csv_keyed(PROCESSED_DIR / "year_actual_10min.csv", 52560, ("date", "slot_id"), validator, "year_actual_10min")
    price_rows = validate_csv_keyed(PROCESSED_DIR / "dynamic_price_10min.csv", 52560, ("date", "slot_id"), validator, "dynamic_price_10min")

    expected_date_text = [value.isoformat() for value in expected_dates()]
    for scope, rows in (("year_load_10min", load_rows), ("year_pv_actual_10min", pv_rows), ("dynamic_price_10min", price_rows)):
        dates = sorted({row["date"] for row in rows})
        per_date = Counter(row["date"] for row in rows)
        validator.check(scope, "365 continuous dates", "2025-01-01..2025-12-31", f"{dates[0]}..{dates[-1]} ({len(dates)})", dates == expected_date_text)
        validator.check(scope, "144 rows per date", "365 dates x 144", f"{len(per_date)} dates; counts={sorted(set(per_date.values()))}", len(per_date) == 365 and set(per_date.values()) == {144})

    load_keys = {(row["date"], row["slot_id"]) for row in load_rows}
    pv_keys = {(row["date"], row["slot_id"]) for row in pv_rows}
    merged_keys = {(row["date"], row["slot_id"]) for row in merged_rows}
    validator.check("year_actual_10min", "no load/PV key lost", 52560, len(load_keys & pv_keys & merged_keys), load_keys == pv_keys == merged_keys and len(merged_keys) == 52560)

    crosswalk = read_rows(TABLES_DIR / "time_marker_crosswalk.csv")
    validator.check("time_crosswalk", "144 slot rows", 144, len(crosswalk), len(crosswalk) == 144)
    validator.check("time_crosswalk", "all-source equality flags", "true", sorted({row["all_sources_same_marker"] for row in crosswalk}), all(row["all_sources_same_marker"] == "true" for row in crosswalk))
    endpoints = (crosswalk[0]["attachment1_time_marker"], crosswalk[142]["attachment1_time_marker"], crosswalk[143]["attachment1_time_marker"])
    validator.check("time_crosswalk", "marker endpoints", ("0:10", "23:50", "0:00+1"), endpoints, endpoints == ("0:10", "23:50", "0:00+1"))

    attachment3_path = ROOT / "input" / "original" / "附件3.xlsx"
    workbook = load_workbook(attachment3_path, data_only=False, read_only=False)
    sheet = workbook["Sheet1"]
    expected_headers = [f"预报{i}小时" for i in range(1, 25)]
    actual_headers = [str(sheet.cell(1, column).value) for column in range(3, 27)]
    validator.check("attachment3", "used range", "A1:Z1461", sheet.calculate_dimension(), sheet.calculate_dimension() == "A1:Z1461")
    validator.check("attachment3", "source issue rows", 1460, sheet.max_row - 1, sheet.max_row - 1 == 1460)
    validator.check("attachment3", "forecast horizon headers", "1..24", len(actual_headers), actual_headers == expected_headers)

    xml_trailing = trailing_cells_ooxml(attachment3_path)
    expected_anchor_dates = expected_dates()
    anchor_dates: list[date | None] = []
    clock_ok_blocks = 0
    semantic_blank_cells = 0
    valid_blocks = 0
    ooxml_index_28_resolves_empty = 0
    for block in range(365):
        start_row = 2 + block * 4
        rows = list(range(start_row, start_row + 4))
        anchor = parse_date(sheet.cell(start_row, 1).value)
        anchor_dates.append(anchor)
        clocks_ok = [normalize_time(sheet.cell(row, 2).value) for row in rows] == ["0:00", "6:00", "12:00", "18:00"]
        clock_ok_blocks += int(clocks_ok)
        trailing_values = [sheet.cell(row, 1).value for row in rows[1:]]
        semantic_ok = all(isinstance(value, str) and value == "" for value in trailing_values)
        semantic_blank_cells += sum(isinstance(value, str) and value == "" for value in trailing_values)
        valid_blocks += int(anchor == expected_anchor_dates[block] and clocks_ok and semantic_ok)
        for row in rows[1:]:
            cell_type, raw_v, resolved = xml_trailing[f"A{row}"]
            ooxml_index_28_resolves_empty += int(cell_type == "s" and raw_v == "28" and resolved == "")
    workbook.close()

    validator.check("attachment3", "candidate four-row block count", 365, len(anchor_dates), len(anchor_dates) == 365)
    validator.check("attachment3", "0h anchor dates continuous and unique", "2025-01-01..2025-12-31 (365 unique)", f"{anchor_dates[0]}..{anchor_dates[-1]} ({len(set(anchor_dates))} unique)", anchor_dates == expected_anchor_dates and len(set(anchor_dates)) == 365)
    validator.check("attachment3", "exact issue-clock sequence blocks", 365, clock_ok_blocks, clock_ok_blocks == 365)
    validator.check("attachment3", "trailing date semantic blanks", 1095, semantic_blank_cells, semantic_blank_cells == 1095)
    validator.check("attachment3", "verified four-row blocks", 365, valid_blocks, valid_blocks == 365)
    validator.check("attachment3", "OOXML shared-string index 28 resolves empty", 1095, ooxml_index_28_resolves_empty, ooxml_index_28_resolves_empty == 1095, blocking=False, detail="Storage evidence only; literal 28 is not a business condition.")

    block_audit = read_rows(TABLES_DIR / "attachment3_block_audit.csv")
    validator.check("attachment3", "block-audit row count", 365, len(block_audit), len(block_audit) == 365)
    validator.check("attachment3", "block-audit verified count", 365, sum(row["structure_ok"] == "true" for row in block_audit), len(block_audit) == 365 and all(row["structure_ok"] == "true" for row in block_audit))

    long_rows = validate_csv_keyed(PROCESSED_DIR / "pv_forecast_hourly_long.csv", 35040, ("issue_datetime", "horizon_hour"), validator, "pv_forecast_hourly_long")
    issue_counts = Counter(row["issue_datetime"] for row in long_rows)
    horizons: dict[str, set[int]] = defaultdict(set)
    clocks_by_date: dict[str, set[str]] = defaultdict(set)
    for row in long_rows:
        horizons[row["issue_datetime"]].add(int(row["horizon_hour"]))
        clocks_by_date[row["issue_date"]].add(row["issue_clock"])
    validator.check("pv_forecast_hourly_long", "unique issue rows", 1460, len(issue_counts), len(issue_counts) == 1460)
    validator.check("pv_forecast_hourly_long", "365 x 4 issue structure", "365 dates x 4 clocks", f"{len(clocks_by_date)} dates; clock-counts={sorted(set(map(len, clocks_by_date.values())))}", len(clocks_by_date) == 365 and all(value == {"0:00", "6:00", "12:00", "18:00"} for value in clocks_by_date.values()))
    complete_horizons = sum(count == 24 and horizons[key] == set(range(1, 25)) for key, count in issue_counts.items())
    validator.check("pv_forecast_hourly_long", "24 horizons and horizon_hour 1..24 per issue", 1460, complete_horizons, complete_horizons == 1460)
    target_formula_count = sum(datetime.fromisoformat(row["nominal_target_datetime"]) == datetime.fromisoformat(row["issue_datetime"]) + timedelta(hours=int(row["horizon_hour"])) for row in long_rows)
    validator.check("pv_forecast_hourly_long", "nominal target formula", 35040, target_formula_count, target_formula_count == 35040)
    cross_year = [row for row in long_rows if datetime.fromisoformat(row["nominal_target_datetime"]).year == 2026]
    maximum_target = max(row["nominal_target_datetime"] for row in long_rows)
    validator.check("pv_forecast_hourly_long", "cross-year nominal targets preserved", "40; max=2026-01-01 18:00:00", f"{len(cross_year)}; max={maximum_target}", len(cross_year) == 40 and maximum_target == "2026-01-01 18:00:00")
    derivations = Counter(row["date_derivation"] for row in long_rows)
    validator.check("pv_forecast_hourly_long", "date derivation labels", {"original_0h_date": 8760, "from_verified_0h_block": 26280}, dict(derivations), derivations == {"original_0h_date": 8760, "from_verified_0h_block": 26280})
    derived_raw_preserved = sum(row["date_derivation"] == "from_verified_0h_block" and row["date_raw"] == "" for row in long_rows)
    validator.check("pv_forecast_hourly_long", "derived-row original semantic blanks preserved", 26280, derived_raw_preserved, derived_raw_preserved == 26280)

    template_contract = read_rows(TABLES_DIR / "template_contract.csv")
    validator.check("templates", "template-contract sheet rows", 16, len(template_contract), len(template_contract) == 16)
    data_issues = read_rows(TABLES_DIR / "data_issues.csv")
    issue_types = Counter(row["issue_type"] for row in data_issues)
    obsolete_blockers = issue_types.get("ATTACHMENT3_TRAILING_DATE_NOT_LITERAL_28", 0) + issue_types.get("BLOCKED_ATTACHMENT3_DATE_DERIVATION", 0)
    validator.check("issues", "obsolete attachment3 blockers absent", 0, obsolete_blockers, obsolete_blockers == 0)

    rewrite_validation_checks(validator.rows)
    if not quiet:
        print(f"VALIDATION STATUS: {'FAIL' if validator.blocking_failures else 'PASS'}")
        print(f"Validator checks: {len(validator.rows)}")
        print(f"Blocking failures: {len(validator.blocking_failures)}")
    return 1 if validator.blocking_failures else 0


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--quiet", action="store_true")
    args = parser.parse_args()
    raise SystemExit(validate(quiet=args.quiet))

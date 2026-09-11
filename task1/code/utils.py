from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import load_workbook


plt.rcParams["font.sans-serif"] = [
    "Noto Sans CJK SC",
    "WenQuanYi Zen Hei",
    "SimHei",
]
plt.rcParams["axes.unicode_minus"] = False
plt.rcParams["savefig.bbox"] = "tight"


def load_problem1_data(path: Path):
    """Load attachment 1 and rotate its cyclic rows to midnight-first order."""
    sheet = load_workbook(path, data_only=True).active
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    if len(rows) != 144 or any(value is None for row in rows for value in row):
        raise ValueError("附件1应包含144行无缺失的10分钟数据")

    # The final 0:00+1 row is the repeating day's 0:00-0:10 interval.
    rows = [rows[-1], *rows[:-1]]
    return {
        "labels": ["0:00", *[str(row[0]) for row in rows[1:]]],
        "price": [float(row[1]) for row in rows],
        "load_power": [float(row[2]) for row in rows],
        "pv_power": [float(row[3]) for row in rows],
    }

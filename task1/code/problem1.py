import csv
import json
import shutil
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import lil_matrix

from utils import load_problem1_data


ROOT = Path(__file__).resolve().parents[1]
INPUT_PATH = ROOT / "附件" / "附件1.xlsx"
TEMPLATE_PATH = ROOT / "附件" / "附件5" / "result1.xlsx"
RESULT_PATH = ROOT / "result1.xlsx"
OUTPUT_DIR = ROOT / "code" / "outputs"
FIGURE_DIR = ROOT / "figures"

N = 144
DT = 1 / 6
ETA_CHARGE = 0.9
ETA_DISCHARGE = 0.9
SOC_INITIAL = 6000.0
SOC_MIN = 1200.0
SOC_MAX = 10800.0
MAX_ENERGY = 5000.0 * DT


def variable_slices():
    starts = np.cumsum([0, N, N, N, N, N + 1])
    return {
        "grid": slice(starts[0], starts[1]),
        "charge": slice(starts[1], starts[2]),
        "discharge": slice(starts[2], starts[3]),
        "spill": slice(starts[3], starts[4]),
        "soc": slice(starts[4], starts[5]),
        "mode": slice(starts[5], starts[5] + N),
    }, starts[5] + N


def solve(data, eta_charge=ETA_CHARGE, eta_discharge=ETA_DISCHARGE):
    idx, size = variable_slices()
    price = np.asarray(data["price"])
    load = np.asarray(data["load_power"]) * DT
    pv = np.asarray(data["pv_power"]) * DT

    objective = np.zeros(size)
    objective[idx["grid"]] = price

    lower = np.zeros(size)
    upper = np.full(size, np.inf)
    upper[idx["charge"]] = MAX_ENERGY
    upper[idx["discharge"]] = MAX_ENERGY
    upper[idx["spill"]] = pv
    lower[idx["soc"]] = SOC_MIN
    upper[idx["soc"]] = SOC_MAX
    lower[idx["soc"].start] = upper[idx["soc"].start] = SOC_INITIAL
    lower[idx["soc"].stop - 1] = upper[idx["soc"].stop - 1] = SOC_INITIAL
    upper[idx["mode"]] = 1

    # Equalities: power balance and storage transition.
    equalities = lil_matrix((2 * N, size))
    rhs = np.zeros(2 * N)
    for t in range(N):
        equalities[t, idx["grid"].start + t] = 1
        equalities[t, idx["charge"].start + t] = -1
        equalities[t, idx["discharge"].start + t] = 1
        equalities[t, idx["spill"].start + t] = -1
        rhs[t] = load[t] - pv[t]

        row = N + t
        equalities[row, idx["soc"].start + t] = -1
        equalities[row, idx["soc"].start + t + 1] = 1
        equalities[row, idx["charge"].start + t] = -eta_charge
        equalities[row, idx["discharge"].start + t] = 1 / eta_discharge

    # Inequalities: charge <= M*z and discharge <= M*(1-z).
    inequalities = lil_matrix((2 * N, size))
    ineq_upper = np.zeros(2 * N)
    for t in range(N):
        inequalities[t, idx["charge"].start + t] = 1
        inequalities[t, idx["mode"].start + t] = -MAX_ENERGY

        row = N + t
        inequalities[row, idx["discharge"].start + t] = 1
        inequalities[row, idx["mode"].start + t] = MAX_ENERGY
        ineq_upper[row] = MAX_ENERGY

    constraints = [
        LinearConstraint(equalities.tocsr(), rhs, rhs),
        LinearConstraint(inequalities.tocsr(), -np.inf, ineq_upper),
    ]
    integrality = np.zeros(size)
    integrality[idx["mode"]] = 1
    result = milp(
        objective,
        integrality=integrality,
        bounds=Bounds(lower, upper),
        constraints=constraints,
        options={"time_limit": 120, "mip_rel_gap": 1e-9},
    )
    if not result.success:
        raise RuntimeError(f"优化失败: {result.message}")

    values = {name: result.x[sl] for name, sl in idx.items()}
    values.update({"load": load, "pv": pv, "objective": float(result.fun)})
    return values


def validate(values):
    grid = values["grid"]
    charge = values["charge"]
    discharge = values["discharge"]
    spill = values["spill"]
    soc = values["soc"]
    balance = grid + values["pv"] + discharge - values["load"] - charge - spill
    transition = soc[1:] - soc[:-1] - ETA_CHARGE * charge + discharge / ETA_DISCHARGE
    baseline_grid = np.maximum(values["load"] - values["pv"], 0)

    return {
        "solver_objective_yuan": values["objective"],
        "total_purchase_kwh": float(grid.sum()),
        "total_charge_kwh": float(charge.sum()),
        "total_discharge_kwh": float(discharge.sum()),
        "total_spill_kwh": float(spill.sum()),
        "baseline_purchase_kwh": float(baseline_grid.sum()),
        "baseline_cost_yuan": None,
        "max_balance_residual_kwh": float(np.max(np.abs(balance))),
        "max_soc_transition_residual_kwh": float(np.max(np.abs(transition))),
        "min_soc_kwh": float(soc.min()),
        "max_soc_kwh": float(soc.max()),
        "initial_soc_kwh": float(soc[0]),
        "final_soc_kwh": float(soc[-1]),
        "max_charge_kwh": float(charge.max()),
        "max_discharge_kwh": float(discharge.max()),
        "simultaneous_charge_discharge_intervals": int(
            np.sum((charge > 1e-6) & (discharge > 1e-6))
        ),
    }


def write_intermediate(data, values, checks):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    csv_path = OUTPUT_DIR / "problem1_schedule.csv"
    with csv_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            [
                "时段起点",
                "电价(元/kWh)",
                "负载电量(kWh)",
                "光伏电量(kWh)",
                "计划购电量(kWh)",
                "充电量(kWh)",
                "放电量(kWh)",
                "弃光量(kWh)",
                "时段末储电量(kWh)",
            ]
        )
        for t in range(N):
            writer.writerow(
                [
                    data["labels"][t],
                    data["price"][t],
                    values["load"][t],
                    values["pv"][t],
                    values["grid"][t],
                    values["charge"][t],
                    values["discharge"][t],
                    values["spill"][t],
                    values["soc"][t + 1],
                ]
            )
    with (OUTPUT_DIR / "problem1_checks.json").open("w", encoding="utf-8") as handle:
        json.dump(checks, handle, ensure_ascii=False, indent=2)


def fill_result_workbook(values):
    shutil.copy2(TEMPLATE_PATH, RESULT_PATH)
    workbook = load_workbook(RESULT_PATH)
    plan_sheet = workbook["计划购电量"]

    # Template order is 0:10,...,23:50,next-day 0:00.
    template_order = np.r_[np.arange(1, N), 0]
    for row, t in enumerate(template_order, start=2):
        plan_sheet.cell(row=row, column=2, value=float(values["grid"][t]))

    storage_sheet = workbook["充放电量"]
    for block in range(6):
        start = block * 24
        stop = (block + 1) * 24
        storage_sheet.cell(block + 2, 2, float(values["charge"][start:stop].sum()))
        storage_sheet.cell(block + 2, 3, float(values["discharge"][start:stop].sum()))
    storage_sheet.cell(2, 5, float(values["soc"][0]))
    storage_sheet.cell(3, 5, float(values["soc"][-1]))
    workbook.save(RESULT_PATH)


def create_figures(data, values):
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    hours = np.arange(N) / 6
    net_storage = values["discharge"] - values["charge"]

    fig, ax = plt.subplots(figsize=(10, 4.8))
    ax.plot(hours, values["load"], label="负载电量", color="#25344F", linewidth=1.8)
    ax.plot(hours, values["pv"], label="光伏电量", color="#E6A23C", linewidth=1.8)
    ax.plot(hours, values["grid"], label="计划购电量", color="#3A7D6F", linewidth=1.5)
    ax.bar(hours, net_storage, width=1 / 6, label="储能净放电量", color="#A45A52", alpha=0.55)
    ax.set_xlabel("时刻/h")
    ax.set_ylabel("每10分钟电量/kWh")
    ax.set_xlim(0, 24)
    ax.grid(axis="y", alpha=0.25)
    ax.legend(ncol=4, loc="upper center")
    fig.savefig(FIGURE_DIR / "problem1_dispatch.pdf")
    plt.close(fig)

    fig, ax1 = plt.subplots(figsize=(10, 4.8))
    ax1.step(np.arange(N + 1) / 6, values["soc"], where="post", color="#25344F", label="储电量")
    ax1.axhline(SOC_MIN, color="#888888", linestyle="--", linewidth=0.9)
    ax1.axhline(SOC_MAX, color="#888888", linestyle="--", linewidth=0.9)
    ax1.set_xlabel("时刻/h")
    ax1.set_ylabel("储电量/kWh")
    ax1.set_xlim(0, 24)
    ax1.grid(axis="y", alpha=0.25)
    ax2 = ax1.twinx()
    ax2.step(hours, data["price"], where="post", color="#C65F3E", alpha=0.8, label="电价")
    ax2.set_ylabel("电价/(元/kWh)")
    lines = ax1.lines[:1] + ax2.lines
    ax1.legend(lines, [line.get_label() for line in lines], loc="upper center", ncol=2)
    fig.savefig(FIGURE_DIR / "problem1_soc_price.pdf")
    plt.close(fig)


def main():
    data = load_problem1_data(INPUT_PATH)
    values = solve(data)
    checks = validate(values)
    baseline_grid = np.maximum(values["load"] - values["pv"], 0)
    checks["baseline_cost_yuan"] = float(np.dot(data["price"], baseline_grid))
    checks["cost_saving_yuan"] = checks["baseline_cost_yuan"] - checks["solver_objective_yuan"]
    checks["cost_saving_rate"] = checks["cost_saving_yuan"] / checks["baseline_cost_yuan"]
    round_trip_split_efficiency = np.sqrt(0.9)
    alternative = solve(data, round_trip_split_efficiency, round_trip_split_efficiency)
    checks["sensitivity_round_trip_90_percent_cost_yuan"] = alternative["objective"]
    checks["sensitivity_cost_change_yuan"] = alternative["objective"] - values["objective"]

    write_intermediate(data, values, checks)
    fill_result_workbook(values)
    create_figures(data, values)

    print(f"最优购电费: {checks['solver_objective_yuan']:.4f} 元")
    print(f"全天购电量: {checks['total_purchase_kwh']:.4f} kWh")
    print(f"相对无储能节省: {checks['cost_saving_yuan']:.4f} 元 ({checks['cost_saving_rate']:.2%})")
    print(f"能量平衡最大残差: {checks['max_balance_residual_kwh']:.3e} kWh")
    print(f"储电量范围: [{checks['min_soc_kwh']:.4f}, {checks['max_soc_kwh']:.4f}] kWh")


if __name__ == "__main__":
    main()

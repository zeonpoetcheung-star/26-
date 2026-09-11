from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from plot_style import configure_plotting


configure_plotting()


def load_problem1_data(path: Path):
    """Load attachment 1 and rotate its cyclic rows to midnight-first order."""
    sheet = load_workbook(path, data_only=True).active
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    if len(rows) != 144 or any(value is None for row in rows for value in row):
        raise ValueError("附件1应包含144行无缺失的10分钟数据")

    # The final 0:00+1 row is the repeating day's 0:00-0:10 interval.
    rows = [rows[-1], *rows[:-1]]
    return {
        "labels": ["0:00", *[str(row[0]) for row in rows[1:]]],
        "price": [float(row[1]) for row in rows],
        "load_power": [float(row[2]) for row in rows],
        "pv_power": [float(row[3]) for row in rows],
    }


def load_problem2_data(path: Path):
    """Load annual load/PV matrices and rotate columns to midnight-first order."""
    workbook = load_workbook(path, data_only=True, read_only=True)
    loaded = {}
    dates = None
    for sheet_name, key in (("小区负载", "load_power"), ("光伏发电实际功率", "pv_power")):
        rows = list(workbook[sheet_name].iter_rows(min_row=2, values_only=True))
        current_dates = [row[0].date() for row in rows]
        values = np.asarray([[float(row[-1]), *map(float, row[1:-1])] for row in rows])
        if values.shape != (365, 144) or not np.isfinite(values).all():
            raise ValueError(f"{sheet_name}应包含365天、每天144个无缺失时段")
        if dates is not None and current_dates != dates:
            raise ValueError("附件2两个工作表的日期不一致")
        dates = current_dates
        loaded[key] = values
    loaded["dates"] = dates
    return loaded

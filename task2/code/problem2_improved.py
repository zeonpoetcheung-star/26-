from __future__ import annotations

import csv
import json
import shutil
from copy import copy
from datetime import datetime
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook

import problem2 as base
from plot_style import save_high_res


ROOT = Path(__file__).resolve().parents[1]
SOURCE_SCHEDULE = ROOT / "code" / "outputs" / "problem2_schedule.csv"
SOURCE_DAILY = ROOT / "code" / "outputs" / "problem2_daily_summary.csv"
SOURCE_CHECKS = ROOT / "code" / "outputs" / "problem2_checks.json"
TEMPLATE_PATH = ROOT / "附件" / "附件5" / "result2.xlsx"
RESULT_PATH = ROOT / "result2_improved.xlsx"
OUTPUT_DIR = ROOT / "code" / "outputs"
FIGURE_DIR = ROOT / "figures"

N = 144
DT = 1 / 6
LOW_QUANTILE = 0.75
HIGH_QUANTILE = 0.85


def load_cached_forecasts():
    rows = list(csv.DictReader(SOURCE_SCHEDULE.open(encoding="utf-8-sig")))
    dates = []
    grouped = {}
    for row in rows:
        date = row["日期"]
        if date not in grouped:
            dates.append(date)
            grouped[date] = []
        grouped[date].append(row)
    if len(dates) != 334 or any(len(grouped[date]) != N for date in dates):
        raise ValueError("原问题二时段结果应包含334天、每天144个时段")

    def matrix(column):
        return np.asarray([[float(row[column]) for row in grouped[date]] for date in dates])

    daily_rows = list(csv.DictReader(SOURCE_DAILY.open(encoding="utf-8-sig")))
    initial_soc = float(daily_rows[0]["0:00储电量(kWh)"])
    return {
        "dates": [datetime.fromisoformat(date).date() for date in dates],
        "price": matrix("电价")[0],
        "load": matrix("实际负载电量"),
        "pv": matrix("实际光伏电量"),
        "load_forecast": matrix("预测负载电量"),
        "pv_forecast": matrix("预测光伏电量"),
        "old_margin": matrix("安全裕度电量"),
        "initial_soc": initial_soc,
        "old_daily": daily_rows,
    }


def adaptive_quantile(soc):
    empty_fraction = (base.SOC_MAX - soc) / (base.SOC_MAX - base.SOC_MIN)
    return LOW_QUANTILE + (HIGH_QUANTILE - LOW_QUANTILE) * np.clip(empty_fraction, 0, 1)


def adaptive_margin(residual_history, old_margin, quantile):
    if len(residual_history) < 7:
        scale = np.clip((quantile - 0.5) / 0.3, 0, 1.2)
        return old_margin * scale
    history = np.asarray(residual_history[-28:])
    margin = np.zeros(N)
    for slot in range(N):
        start = max(0, slot - 3)
        stop = min(N, slot + 4)
        margin[slot] = max(0.0, float(np.quantile(history[:, start:stop], quantile)))
    return margin


def operate_causal(grid, actual_load, actual_pv, initial_soc):
    charge = np.zeros(N)
    discharge = np.zeros(N)
    emergency = np.zeros(N)
    spill = np.zeros(N)
    soc = np.zeros(N + 1)
    soc[0] = initial_soc
    for slot in range(N):
        balance = grid[slot] + actual_pv[slot] - actual_load[slot]
        if balance >= 0:
            charge[slot] = min(
                balance,
                base.MAX_ENERGY,
                (base.SOC_MAX - soc[slot]) / base.ETA_CHARGE,
            )
            spill[slot] = balance - charge[slot]
            soc[slot + 1] = soc[slot] + base.ETA_CHARGE * charge[slot]
            continue

        deficit = -balance
        usable_output = (soc[slot] - base.SOC_MIN) * base.ETA_DISCHARGE
        discharge[slot] = min(deficit, base.MAX_ENERGY, usable_output)
        emergency[slot] = deficit - discharge[slot]
        soc[slot + 1] = soc[slot] - discharge[slot] / base.ETA_DISCHARGE

    return {
        "charge": charge,
        "discharge": discharge,
        "emergency": emergency,
        "spill": spill,
        "soc": soc,
    }


def run_improved(data):
    days = len(data["dates"])
    outputs = {key: np.zeros((days, N)) for key in ("grid", "charge", "discharge", "emergency", "spill", "margin")}
    outputs["soc"] = np.zeros((days, N + 1))
    outputs["quantile"] = np.zeros(days)
    residual_history = []
    soc = data["initial_soc"]

    for day in range(days):
        quantile = adaptive_quantile(soc)
        margin = adaptive_margin(residual_history, data["old_margin"][day], quantile)
        plan = base.plan_day(
            data["price"],
            (data["load_forecast"][day] + margin) / DT,
            data["pv_forecast"][day] / DT,
            soc,
        )
        actual = operate_causal(
            plan["grid"],
            data["load"][day],
            data["pv"][day],
            soc,
        )
        outputs["grid"][day] = plan["grid"]
        outputs["margin"][day] = margin
        outputs["quantile"][day] = quantile
        for key in ("charge", "discharge", "emergency", "spill"):
            outputs[key][day] = actual[key]
        outputs["soc"][day] = actual["soc"]
        soc = actual["soc"][-1]
        residual_history.append(
            (data["load"][day] - data["pv"][day])
            - (data["load_forecast"][day] - data["pv_forecast"][day])
        )

    outputs["plan_cost"] = outputs["grid"] @ data["price"]
    outputs["emergency_cost"] = 5 * (outputs["emergency"] @ data["price"])
    outputs["total_cost"] = outputs["plan_cost"] + outputs["emergency_cost"]
    return outputs


def apply_style(sheet, styles, row):
    for column, style in enumerate(styles, start=1):
        sheet.cell(row, column)._style = copy(style)


def write_workbook(data, outputs):
    shutil.copy2(TEMPLATE_PATH, RESULT_PATH)
    workbook = load_workbook(RESULT_PATH)
    order = np.r_[np.arange(1, N), 0]
    plan_sheet = workbook["计划购电量"]
    for row, day in enumerate(range(len(data["dates"])), start=2):
        plan_sheet.cell(row, 1, datetime.combine(data["dates"][day], datetime.min.time()))
        for column, slot in enumerate(order, start=2):
            plan_sheet.cell(row, column, float(outputs["grid"][day, slot]))
        plan_sheet.cell(row, 146, float(outputs["grid"][day].sum()))
        plan_sheet.cell(row, 147, float(outputs["plan_cost"][day]))

    storage = workbook["充放电量"]
    styles = [[copy(storage.cell(row, column)._style) for column in range(1, 7)] for row in range(2, 8)]
    storage.delete_rows(2, storage.max_row - 1)
    blocks = ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]
    row = 2
    for day, date in enumerate(data["dates"]):
        for block, label in enumerate(blocks):
            apply_style(storage, styles[block], row)
            start, stop = block * 24, (block + 1) * 24
            storage.cell(row, 1, datetime.combine(date, datetime.min.time()) if block == 0 else None)
            storage.cell(row, 2, label)
            storage.cell(row, 3, float(outputs["charge"][day, start:stop].sum()))
            storage.cell(row, 4, float(outputs["discharge"][day, start:stop].sum()))
            if block == 0:
                storage.cell(row, 5, "0:00")
                storage.cell(row, 6, float(outputs["soc"][day, 0]))
            elif block == 1:
                storage.cell(row, 5, "24:00")
                storage.cell(row, 6, float(outputs["soc"][day, -1]))
            row += 1

    emergency = workbook["紧急购电量"]
    style = [copy(emergency.cell(2, column)._style) for column in range(1, 4)]
    emergency.delete_rows(2, emergency.max_row - 1)
    row = 2
    for day, date in enumerate(data["dates"]):
        for index, (period, amount) in enumerate(base.emergency_segments(outputs["emergency"][day])):
            apply_style(emergency, style, row)
            emergency.cell(row, 1, datetime.combine(date, datetime.min.time()) if index == 0 else None)
            emergency.cell(row, 2, period)
            emergency.cell(row, 3, amount)
            row += 1
    workbook.save(RESULT_PATH)


def write_outputs(data, outputs):
    schedule_path = OUTPUT_DIR / "problem2_improved_schedule.csv"
    with schedule_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(["日期", "时段", "电价", "实际负载电量", "实际光伏电量", "预测负载电量", "预测光伏电量", "自适应分位数", "安全裕度电量", "计划购电量", "实际充电量", "实际放电量", "紧急购电量", "弃电量", "时段末储电量"])
        for day, date in enumerate(data["dates"]):
            for slot in range(N):
                writer.writerow([date.isoformat(), base.format_clock(slot * 10), data["price"][slot], data["load"][day, slot], data["pv"][day, slot], data["load_forecast"][day, slot], data["pv_forecast"][day, slot], outputs["quantile"][day], outputs["margin"][day, slot], outputs["grid"][day, slot], outputs["charge"][day, slot], outputs["discharge"][day, slot], outputs["emergency"][day, slot], outputs["spill"][day, slot], outputs["soc"][day, slot + 1]])

    daily_path = OUTPUT_DIR / "problem2_improved_daily_summary.csv"
    with daily_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(["日期", "自适应分位数", "计划购电量(kWh)", "紧急购电量(kWh)", "计划购电费(元)", "紧急购电费(元)", "总费用(元)", "0:00储电量(kWh)", "24:00储电量(kWh)"])
        for day, date in enumerate(data["dates"]):
            writer.writerow([date.isoformat(), outputs["quantile"][day], outputs["grid"][day].sum(), outputs["emergency"][day].sum(), outputs["plan_cost"][day], outputs["emergency_cost"][day], outputs["total_cost"][day], outputs["soc"][day, 0], outputs["soc"][day, -1]])

    with SOURCE_CHECKS.open(encoding="utf-8") as handle:
        old = json.load(handle)
    old_total = old["total_cost_yuan"]
    balance = outputs["grid"] + data["pv"] + outputs["discharge"] + outputs["emergency"] - data["load"] - outputs["charge"] - outputs["spill"]
    transition = outputs["soc"][:, 1:] - outputs["soc"][:, :-1] - base.ETA_CHARGE * outputs["charge"] + outputs["discharge"] / base.ETA_DISCHARGE
    checks = {
        "method": "SOC-adaptive residual quantile plus causal deficit-first dispatch",
        "days": len(data["dates"]),
        "initial_soc_kwh": data["initial_soc"],
        "average_quantile": float(outputs["quantile"].mean()),
        "min_quantile": float(outputs["quantile"].min()),
        "max_quantile": float(outputs["quantile"].max()),
        "plan_purchase_kwh": float(outputs["grid"].sum()),
        "emergency_purchase_kwh": float(outputs["emergency"].sum()),
        "plan_cost_yuan": float(outputs["plan_cost"].sum()),
        "emergency_cost_yuan": float(outputs["emergency_cost"].sum()),
        "total_cost_yuan": float(outputs["total_cost"].sum()),
        "old_total_cost_yuan": old_total,
        "cost_change_yuan": float(outputs["total_cost"].sum() - old_total),
        "cost_change_rate": float(outputs["total_cost"].sum() / old_total - 1),
        "emergency_intervals": int(np.sum(outputs["emergency"] > 1e-6)),
        "emergency_segments": int(sum(len(base.emergency_segments(day)) for day in outputs["emergency"])),
        "days_starting_at_soc_max": int(np.sum(np.isclose(outputs["soc"][:, 0], base.SOC_MAX))),
        "min_soc_kwh": float(outputs["soc"].min()),
        "max_soc_kwh": float(outputs["soc"].max()),
        "max_charge_kwh": float(outputs["charge"].max()),
        "max_discharge_kwh": float(outputs["discharge"].max()),
        "max_balance_residual_kwh": float(np.max(np.abs(balance))),
        "max_transition_residual_kwh": float(np.max(np.abs(transition))),
        "cross_day_soc_residual_kwh": float(np.max(np.abs(outputs["soc"][:-1, -1] - outputs["soc"][1:, 0]))),
    }
    workbook = load_workbook(RESULT_PATH, read_only=True, data_only=True)
    plan_rows = list(workbook["计划购电量"].iter_rows(min_row=2, values_only=True))
    workbook_plan = sum(sum(float(value or 0) for value in row[1:145]) for row in plan_rows)
    workbook_cost = sum(float(row[146] or 0) for row in plan_rows)
    emergency_rows = list(workbook["紧急购电量"].iter_rows(min_row=2, values_only=True))
    workbook_emergency = sum(float(row[2] or 0) for row in emergency_rows)
    checks.update(
        {
            "workbook_plan_rows": len(plan_rows),
            "workbook_storage_rows": workbook["充放电量"].max_row - 1,
            "workbook_emergency_segments": len(emergency_rows),
            "workbook_plan_residual_kwh": float(workbook_plan - outputs["grid"].sum()),
            "workbook_cost_residual_yuan": float(workbook_cost - outputs["plan_cost"].sum()),
            "workbook_emergency_residual_kwh": float(workbook_emergency - outputs["emergency"].sum()),
        }
    )
    with (OUTPUT_DIR / "problem2_improved_checks.json").open("w", encoding="utf-8") as handle:
        json.dump(checks, handle, ensure_ascii=False, indent=2)
    return checks


def create_figures(data, outputs):
    old_cost = np.asarray([float(row["总费用(元)"]) for row in data["old_daily"]])
    x = np.arange(len(data["dates"]))
    fig, (cost_axis, quantile_axis) = plt.subplots(2, 1, figsize=(11, 6.4), sharex=True)
    cost_axis.plot(x, old_cost, color="#7B8794", linewidth=0.9, label="原方案")
    cost_axis.plot(x, outputs["total_cost"], color="#2A7F8E", linewidth=0.9, label="改进方案")
    cost_axis.set_ylabel("每日总费用 / 元")
    cost_axis.legend(ncol=2)
    quantile_axis.plot(x, outputs["quantile"], color="#B44A3A", linewidth=1.0, label="自适应分位数")
    quantile_axis.set_ylim(0.6, 0.9)
    quantile_axis.set_xlabel("自 2025-02-01 起的日序")
    quantile_axis.set_ylabel("风险分位数")
    quantile_axis.legend()
    fig.tight_layout()
    save_high_res(fig, FIGURE_DIR / "problem2_improved_cost_quantile", formats=("pdf", "png"))

    fig, axis = plt.subplots(figsize=(10.5, 4.8))
    axis.plot(x, outputs["soc"][:, 0], color="#31587A", linewidth=0.9, label="0:00 储电量")
    axis.plot(x, outputs["soc"][:, -1], color="#2A7F8E", linewidth=0.9, label="24:00 储电量")
    axis.axhline(base.SOC_MIN, color="#7B8794", linestyle="--", linewidth=0.8)
    axis.axhline(base.SOC_MAX, color="#7B8794", linestyle="--", linewidth=0.8)
    axis.set_xlabel("自 2025-02-01 起的日序")
    axis.set_ylabel("储电量 / kWh")
    axis.legend(ncol=2)
    save_high_res(fig, FIGURE_DIR / "problem2_improved_soc", formats=("pdf", "png"))


def main():
    data = load_cached_forecasts()
    outputs = run_improved(data)
    write_workbook(data, outputs)
    checks = write_outputs(data, outputs)
    create_figures(data, outputs)
    print(f"改进版计划购电量: {checks['plan_purchase_kwh']:.2f} kWh")
    print(f"改进版紧急购电量: {checks['emergency_purchase_kwh']:.2f} kWh")
    print(f"改进版总费用: {checks['total_cost_yuan']:.2f} 元")
    print(f"相对原方案费用变化: {checks['cost_change_yuan']:.2f} 元 ({checks['cost_change_rate']:.2%})")
    print(f"储电量范围: [{checks['min_soc_kwh']:.2f}, {checks['max_soc_kwh']:.2f}] kWh")


if __name__ == "__main__":
    main()

from __future__ import annotations

import csv
import json
import shutil
from copy import copy
from datetime import datetime
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from scipy.optimize import linprog
from scipy.sparse import lil_matrix
from sklearn.ensemble import RandomForestRegressor
from sklearn.linear_model import Ridge
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler

from plot_style import save_high_res
from utils import load_problem1_data, load_problem2_data


ROOT = Path(__file__).resolve().parents[1]
PRICE_PATH = ROOT / "附件" / "附件1.xlsx"
DATA_PATH = ROOT / "附件" / "附件2.xlsx"
TEMPLATE_PATH = ROOT / "附件" / "附件5" / "result2.xlsx"
RESULT_PATH = ROOT / "result2.xlsx"
OUTPUT_DIR = ROOT / "code" / "outputs"
FIGURE_DIR = ROOT / "figures"

N = 144
DT = 1 / 6
ETA_CHARGE = 0.9
ETA_DISCHARGE = 0.9
SOC_INITIAL = 6000.0
SOC_MIN = 1200.0
SOC_MAX = 10800.0
MAX_ENERGY = 5000.0 * DT
OUTPUT_START = 31  # 2025-02-01
RISK_QUANTILE = 0.8


def day_features(series: np.ndarray, day: int, dates) -> np.ndarray:
    previous = series[day - 1]
    previous_week = series[day - 7]
    recent_mean = series[max(0, day - 7) : day].mean(axis=0)
    date = dates[day]
    angle = 2 * np.pi * date.timetuple().tm_yday / 365
    weekday = np.eye(7)[date.weekday()]
    calendar = np.r_[np.sin(angle), np.cos(angle), weekday, date.weekday() >= 5]
    return np.r_[previous, previous_week, recent_mean, calendar]


def weekday_mean(series: np.ndarray, day: int, dates) -> np.ndarray:
    indices = [
        i
        for i in range(max(0, day - 35), day)
        if dates[i].weekday() == dates[day].weekday()
    ]
    if not indices:
        return series[max(0, day - 7) : day].mean(axis=0)
    return series[indices[-4:]].mean(axis=0)


def inverse_error_weights(error_history: dict[str, list[float]], names: list[str]) -> np.ndarray:
    if min((len(error_history[name]) for name in names), default=0) < 7:
        return np.full(len(names), 1 / len(names))
    recent = np.asarray([np.mean(error_history[name][-28:]) for name in names])
    inverse = 1 / np.maximum(recent, 1e-6)
    return inverse / inverse.sum()


def generate_forecasts(actual: np.ndarray, dates, typical: np.ndarray, seed: int):
    days = actual.shape[0]
    names = ["seasonal", "weekday", "ridge", "forest"]
    predictions = {name: np.zeros_like(actual) for name in names}
    ensemble = np.zeros_like(actual)
    weights = np.zeros((days, len(names)))
    errors = {name: [] for name in names}
    ridge = forest = None

    for day in range(days):
        if day < 7:
            candidates = {name: typical.copy() for name in names}
            active_names = ["seasonal"]
        else:
            candidates = {
                "seasonal": actual[day - 7].copy(),
                "weekday": weekday_mean(actual, day, dates),
            }
            if day >= 21 and (ridge is None or day % 7 == 0 or day == OUTPUT_START):
                train_days = range(7, day)
                x_train = np.asarray([day_features(actual, item, dates) for item in train_days])
                y_train = actual[7:day]
                ridge = make_pipeline(StandardScaler(), Ridge(alpha=30.0))
                ridge.fit(x_train, y_train)
                forest = RandomForestRegressor(
                    n_estimators=100,
                    min_samples_leaf=2,
                    max_features=0.65,
                    n_jobs=-1,
                    random_state=seed,
                )
                forest.fit(x_train, y_train)
            if ridge is None:
                candidates["ridge"] = candidates["weekday"].copy()
                candidates["forest"] = candidates["weekday"].copy()
                active_names = ["seasonal", "weekday"]
            else:
                features = day_features(actual, day, dates).reshape(1, -1)
                candidates["ridge"] = ridge.predict(features)[0]
                candidates["forest"] = forest.predict(features)[0]
                active_names = names

        for name in names:
            values = np.maximum(candidates[name], 0)
            predictions[name][day] = values

        active_weights = inverse_error_weights(errors, active_names)
        weight_map = dict(zip(active_names, active_weights))
        weights[day] = [weight_map.get(name, 0.0) for name in names]
        ensemble[day] = sum(weights[day, i] * predictions[name][day] for i, name in enumerate(names))

        for name in names:
            errors[name].append(float(np.mean(np.abs(predictions[name][day] - actual[day]))))

    return {"candidates": predictions, "ensemble": ensemble, "weights": weights}


def enforce_pv_physics(forecast: np.ndarray, actual: np.ndarray) -> np.ndarray:
    result = np.maximum(forecast, 0)
    for day in range(len(result)):
        history = actual[max(0, day - 28) : day]
        if len(history):
            result[day, np.max(history, axis=0) < 1.0] = 0
    return result


def residual_margins(actual_net: np.ndarray, forecast_net: np.ndarray) -> np.ndarray:
    residual = actual_net - forecast_net
    margins = np.zeros_like(residual)
    for day in range(len(residual)):
        history = residual[max(0, day - 28) : day]
        if not len(history):
            continue
        for slot in range(N):
            start = max(0, slot - 3)
            stop = min(N, slot + 4)
            margins[day, slot] = max(
                0.0,
                float(np.quantile(history[:, start:stop], RISK_QUANTILE)),
            )
    return margins


def variable_slices():
    starts = np.cumsum([0, N, N, N, N, N + 1])
    return {
        "grid": slice(starts[0], starts[1]),
        "charge": slice(starts[1], starts[2]),
        "discharge": slice(starts[2], starts[3]),
        "spill": slice(starts[3], starts[4]),
        "soc": slice(starts[4], starts[5]),
    }, starts[5]


def plan_day(price, load_power, pv_power, initial_soc):
    idx, size = variable_slices()
    load = np.maximum(load_power, 0) * DT
    pv = np.maximum(pv_power, 0) * DT
    objective = np.zeros(size)
    objective[idx["grid"]] = price
    objective[idx["charge"]] = 1e-7
    objective[idx["discharge"]] = 1e-7

    lower = np.zeros(size)
    upper = np.full(size, np.inf)
    upper[idx["charge"]] = MAX_ENERGY
    upper[idx["discharge"]] = MAX_ENERGY
    upper[idx["spill"]] = pv
    lower[idx["soc"]] = SOC_MIN
    upper[idx["soc"]] = SOC_MAX
    lower[idx["soc"].start] = upper[idx["soc"].start] = initial_soc
    lower[idx["soc"].stop - 1] = upper[idx["soc"].stop - 1] = initial_soc

    equalities = lil_matrix((2 * N, size))
    rhs = np.zeros(2 * N)
    for slot in range(N):
        equalities[slot, idx["grid"].start + slot] = 1
        equalities[slot, idx["charge"].start + slot] = -1
        equalities[slot, idx["discharge"].start + slot] = 1
        equalities[slot, idx["spill"].start + slot] = -1
        rhs[slot] = load[slot] - pv[slot]

        row = N + slot
        equalities[row, idx["soc"].start + slot] = -1
        equalities[row, idx["soc"].start + slot + 1] = 1
        equalities[row, idx["charge"].start + slot] = -ETA_CHARGE
        equalities[row, idx["discharge"].start + slot] = 1 / ETA_DISCHARGE

    result = linprog(
        objective,
        A_eq=equalities.tocsr(),
        b_eq=rhs,
        bounds=list(zip(lower, upper)),
        method="highs",
    )
    if not result.success:
        raise RuntimeError(f"全天计划优化失败: {result.message}")
    return {name: result.x[section] for name, section in idx.items()}


def operate_day(grid, actual_load_power, actual_pv_power, initial_soc):
    load = actual_load_power * DT
    pv = actual_pv_power * DT
    charge = np.zeros(N)
    discharge = np.zeros(N)
    emergency = np.zeros(N)
    spill = np.zeros(N)
    soc = np.zeros(N + 1)
    soc[0] = initial_soc

    for slot in range(N):
        balance = grid[slot] + pv[slot] - load[slot]
        if balance >= 0:
            charge[slot] = min(balance, MAX_ENERGY, (SOC_MAX - soc[slot]) / ETA_CHARGE)
            spill[slot] = balance - charge[slot]
            soc[slot + 1] = soc[slot] + ETA_CHARGE * charge[slot]
        else:
            deficit = -balance
            discharge[slot] = min(deficit, MAX_ENERGY, (soc[slot] - SOC_MIN) * ETA_DISCHARGE)
            emergency[slot] = deficit - discharge[slot]
            soc[slot + 1] = soc[slot] - discharge[slot] / ETA_DISCHARGE

    return {
        "charge": charge,
        "discharge": discharge,
        "emergency": emergency,
        "spill": spill,
        "soc": soc,
        "load": load,
        "pv": pv,
    }


def run_strategy(price, load, pv, load_forecast, pv_forecast, margins, start_soc=SOC_INITIAL):
    days = len(load)
    keys = ["grid", "charge", "discharge", "emergency", "spill"]
    outputs = {key: np.zeros((days, N)) for key in keys}
    outputs["soc"] = np.zeros((days, N + 1))
    soc = start_soc
    for day in range(days):
        risk_load = load_forecast[day] + margins[day]
        plan = plan_day(price, risk_load, pv_forecast[day], soc)
        actual = operate_day(plan["grid"], load[day], pv[day], soc)
        outputs["grid"][day] = plan["grid"]
        for key in ("charge", "discharge", "emergency", "spill"):
            outputs[key][day] = actual[key]
        outputs["soc"][day] = actual["soc"]
        soc = actual["soc"][-1]
    outputs["plan_cost"] = outputs["grid"] @ price
    outputs["emergency_cost"] = 5 * (outputs["emergency"] @ price)
    outputs["total_cost"] = outputs["plan_cost"] + outputs["emergency_cost"]
    return outputs


def prediction_metrics(actual, forecast, start=OUTPUT_START):
    error = forecast[start:] - actual[start:]
    return {
        "mae": float(np.mean(np.abs(error))),
        "rmse": float(np.sqrt(np.mean(error**2))),
    }


def format_clock(minutes: int) -> str:
    if minutes == 24 * 60:
        return "24:00"
    return f"{minutes // 60}:{minutes % 60:02d}"


def emergency_segments(values: np.ndarray):
    segments = []
    slot = 0
    while slot < N:
        if values[slot] <= 1e-6:
            slot += 1
            continue
        start = slot
        while slot < N and values[slot] > 1e-6:
            slot += 1
        segments.append((f"{format_clock(start * 10)}-{format_clock(slot * 10)}", float(values[start:slot].sum())))
    return segments


def fill_result_workbook(dates, price, outputs):
    shutil.copy2(TEMPLATE_PATH, RESULT_PATH)
    workbook = load_workbook(RESULT_PATH)
    template_order = np.r_[np.arange(1, N), 0]

    plan_sheet = workbook["计划购电量"]
    for output_row, day in enumerate(range(OUTPUT_START, len(dates)), start=2):
        plan_sheet.cell(output_row, 1, datetime.combine(dates[day], datetime.min.time()))
        for column, slot in enumerate(template_order, start=2):
            plan_sheet.cell(output_row, column, float(outputs["grid"][day, slot]))
        plan_sheet.cell(output_row, 146, float(outputs["grid"][day].sum()))
        plan_sheet.cell(output_row, 147, float(outputs["plan_cost"][day]))

    storage_sheet = workbook["充放电量"]
    storage_styles = [
        [copy(storage_sheet.cell(row, column)._style) for column in range(1, 7)]
        for row in range(2, 8)
    ]
    storage_sheet.delete_rows(2, storage_sheet.max_row - 1)
    output_row = 2
    blocks = ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]
    for day in range(OUTPUT_START, len(dates)):
        for block, label in enumerate(blocks):
            for column in range(1, 7):
                storage_sheet.cell(output_row, column)._style = copy(storage_styles[block][column - 1])
            start, stop = block * 24, (block + 1) * 24
            storage_sheet.cell(output_row, 1, datetime.combine(dates[day], datetime.min.time()) if block == 0 else None)
            storage_sheet.cell(output_row, 2, label)
            storage_sheet.cell(output_row, 3, float(outputs["charge"][day, start:stop].sum()))
            storage_sheet.cell(output_row, 4, float(outputs["discharge"][day, start:stop].sum()))
            if block == 0:
                storage_sheet.cell(output_row, 5, "0:00")
                storage_sheet.cell(output_row, 6, float(outputs["soc"][day, 0]))
            elif block == 1:
                storage_sheet.cell(output_row, 5, "24:00")
                storage_sheet.cell(output_row, 6, float(outputs["soc"][day, -1]))
            output_row += 1

    emergency_sheet = workbook["紧急购电量"]
    emergency_style = [copy(emergency_sheet.cell(2, column)._style) for column in range(1, 4)]
    emergency_sheet.delete_rows(2, emergency_sheet.max_row - 1)
    output_row = 2
    for day in range(OUTPUT_START, len(dates)):
        segments = emergency_segments(outputs["emergency"][day])
        for segment_index, (period, amount) in enumerate(segments):
            for column in range(1, 4):
                emergency_sheet.cell(output_row, column)._style = copy(emergency_style[column - 1])
            emergency_sheet.cell(output_row, 1, datetime.combine(dates[day], datetime.min.time()) if segment_index == 0 else None)
            emergency_sheet.cell(output_row, 2, period)
            emergency_sheet.cell(output_row, 3, amount)
            output_row += 1

    workbook.save(RESULT_PATH)


def validate_result_workbook(dates, price, outputs):
    workbook = load_workbook(RESULT_PATH, read_only=True, data_only=True)
    plan_sheet = workbook["计划购电量"]
    template_order = np.r_[np.arange(1, N), 0]
    excel_grid = np.asarray(
        [
            [float(plan_sheet.cell(row, column).value) for column in range(2, 146)]
            for row in range(2, plan_sheet.max_row + 1)
        ]
    )
    expected_grid = outputs["grid"][OUTPUT_START:][:, template_order]
    storage_sheet = workbook["充放电量"]
    emergency_sheet = workbook["紧急购电量"]
    emergency_sum = sum(
        float(emergency_sheet.cell(row, 3).value or 0)
        for row in range(2, emergency_sheet.max_row + 1)
    )
    return {
        "workbook_plan_days": plan_sheet.max_row - 1,
        "workbook_plan_intervals": int(excel_grid.size),
        "workbook_storage_rows": storage_sheet.max_row - 1,
        "workbook_emergency_segments": emergency_sheet.max_row - 1,
        "workbook_grid_max_difference_kwh": float(np.max(np.abs(excel_grid - expected_grid))),
        "workbook_plan_total_difference_kwh": float(
            abs(excel_grid.sum() - outputs["grid"][OUTPUT_START:].sum())
        ),
        "workbook_emergency_total_difference_kwh": float(
            abs(emergency_sum - outputs["emergency"][OUTPUT_START:].sum())
        ),
        "workbook_dates_match": all(
            plan_sheet.cell(row, 1).value.date() == dates[OUTPUT_START + row - 2]
            for row in range(2, plan_sheet.max_row + 1)
        ),
    }


def write_outputs(dates, price, load, pv, load_forecast, pv_forecast, margins, outputs, metrics, comparisons):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    with (OUTPUT_DIR / "problem2_schedule.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(["日期", "时段", "电价", "实际负载电量", "实际光伏电量", "预测负载电量", "预测光伏电量", "安全裕度电量", "计划购电量", "实际充电量", "实际放电量", "紧急购电量", "弃光量", "时段末储电量"])
        for day in range(OUTPUT_START, len(dates)):
            for slot in range(N):
                writer.writerow([
                    dates[day].isoformat(), format_clock(slot * 10), price[slot], load[day, slot] * DT,
                    pv[day, slot] * DT, load_forecast[day, slot] * DT, pv_forecast[day, slot] * DT,
                    margins[day, slot] * DT, outputs["grid"][day, slot], outputs["charge"][day, slot],
                    outputs["discharge"][day, slot], outputs["emergency"][day, slot], outputs["spill"][day, slot],
                    outputs["soc"][day, slot + 1],
                ])

    with (OUTPUT_DIR / "problem2_daily_summary.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(["日期", "计划购电量(kWh)", "紧急购电量(kWh)", "计划购电费(元)", "紧急购电费(元)", "总费用(元)", "0:00储电量(kWh)", "24:00储电量(kWh)"])
        for day in range(OUTPUT_START, len(dates)):
            writer.writerow([dates[day].isoformat(), outputs["grid"][day].sum(), outputs["emergency"][day].sum(), outputs["plan_cost"][day], outputs["emergency_cost"][day], outputs["total_cost"][day], outputs["soc"][day, 0], outputs["soc"][day, -1]])

    checks = {
        "prediction_metrics_power_kw": metrics,
        "strategy_comparison": comparisons,
        "output_days": len(dates) - OUTPUT_START,
        "plan_intervals": int(np.prod(outputs["grid"][OUTPUT_START:].shape)),
        "total_plan_purchase_kwh": float(outputs["grid"][OUTPUT_START:].sum()),
        "total_emergency_purchase_kwh": float(outputs["emergency"][OUTPUT_START:].sum()),
        "total_plan_cost_yuan": float(outputs["plan_cost"][OUTPUT_START:].sum()),
        "total_emergency_cost_yuan": float(outputs["emergency_cost"][OUTPUT_START:].sum()),
        "total_cost_yuan": float(outputs["total_cost"][OUTPUT_START:].sum()),
        "emergency_intervals": int(np.sum(outputs["emergency"][OUTPUT_START:] > 1e-6)),
        "min_soc_kwh": float(outputs["soc"][OUTPUT_START:].min()),
        "max_soc_kwh": float(outputs["soc"][OUTPUT_START:].max()),
        "max_charge_kwh": float(outputs["charge"][OUTPUT_START:].max()),
        "max_discharge_kwh": float(outputs["discharge"][OUTPUT_START:].max()),
        "cross_day_soc_residual_kwh": float(np.max(np.abs(outputs["soc"][OUTPUT_START:-1, -1] - outputs["soc"][OUTPUT_START + 1 :, 0]))),
    }
    balance = outputs["grid"] + pv * DT + outputs["discharge"] + outputs["emergency"] - load * DT - outputs["charge"] - outputs["spill"]
    transition = outputs["soc"][:, 1:] - outputs["soc"][:, :-1] - ETA_CHARGE * outputs["charge"] + outputs["discharge"] / ETA_DISCHARGE
    checks["max_actual_balance_residual_kwh"] = float(np.max(np.abs(balance[OUTPUT_START:])))
    checks["max_soc_transition_residual_kwh"] = float(np.max(np.abs(transition[OUTPUT_START:])))
    checks.update(validate_result_workbook(dates, price, outputs))
    with (OUTPUT_DIR / "problem2_checks.json").open("w", encoding="utf-8") as handle:
        json.dump(checks, handle, ensure_ascii=False, indent=2)
    return checks


def create_figures(dates, load, pv, load_forecast, pv_forecast, outputs, metrics):
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    model_names = list(metrics["net"].keys())
    model_labels = {"seasonal": "周同期", "weekday": "同星期均值", "ridge": "岭回归", "forest": "随机森林", "ensemble": "自适应集成"}
    x = np.arange(len(model_names))
    fig, ax = plt.subplots(figsize=(8.8, 4.8))
    ax.bar(x - 0.18, [metrics["net"][name]["mae"] for name in model_names], width=0.36, color="#3B6C8E", label="MAE")
    ax.bar(x + 0.18, [metrics["net"][name]["rmse"] for name in model_names], width=0.36, color="#D09B43", label="RMSE")
    ax.set_xticks(x, [model_labels[name] for name in model_names])
    ax.set_ylabel("净负荷预测误差 / kW")
    ax.legend(ncol=2)
    save_high_res(fig, FIGURE_DIR / "problem2_forecast_metrics", formats=("pdf", "png"))

    selected_dates = {"2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"}
    selected = [i for i, date in enumerate(dates) if date.isoformat() in selected_dates]
    fig, axes = plt.subplots(2, 2, figsize=(11.5, 7.0), sharex=True)
    hours = np.arange(N) / 6
    for axis, day in zip(axes.flat, selected):
        actual_net = (load[day] - pv[day]) * DT
        forecast_net = (load_forecast[day] - pv_forecast[day]) * DT
        axis.plot(hours, actual_net, color="#243B53", linewidth=1.6, label="实际净负荷")
        axis.plot(hours, forecast_net, color="#D09B43", linewidth=1.4, label="预测净负荷")
        axis.step(hours, outputs["grid"][day], where="post", color="#2A7F8E", linewidth=1.2, label="计划购电")
        axis.fill_between(hours, outputs["emergency"][day], step="post", color="#C85C5C", alpha=0.55, label="紧急购电")
        axis.set_title(dates[day].isoformat(), fontsize=11)
        axis.set_xlim(0, 24)
        axis.set_ylabel("电量 / kWh")
    axes[-1, 0].set_xlabel("时刻 / h")
    axes[-1, 1].set_xlabel("时刻 / h")
    handles, labels = axes[0, 0].get_legend_handles_labels()
    fig.legend(handles, labels, loc="upper center", ncol=4, bbox_to_anchor=(0.5, 1.01))
    fig.tight_layout(rect=(0, 0, 1, 0.96))
    save_high_res(fig, FIGURE_DIR / "problem2_selected_days", formats=("pdf", "png"))

    day_axis = np.arange(len(dates) - OUTPUT_START)
    fig, (cost_axis, soc_axis) = plt.subplots(2, 1, figsize=(11, 6.2), sharex=True)
    cost_axis.plot(day_axis, outputs["plan_cost"][OUTPUT_START:], color="#31587A", linewidth=1.0, label="计划购电费")
    cost_axis.plot(day_axis, outputs["emergency_cost"][OUTPUT_START:], color="#B44A3A", linewidth=1.0, label="紧急购电费")
    cost_axis.set_ylabel("每日费用 / 元")
    cost_axis.legend(ncol=2)
    soc_axis.fill_between(day_axis, outputs["soc"][OUTPUT_START:, 0], outputs["soc"][OUTPUT_START:, -1], color="#2A7F8E", alpha=0.25)
    soc_axis.plot(day_axis, outputs["soc"][OUTPUT_START:, 0], color="#31587A", linewidth=0.9, label="0:00 储电量")
    soc_axis.plot(day_axis, outputs["soc"][OUTPUT_START:, -1], color="#2A7F8E", linewidth=0.9, label="24:00 储电量")
    soc_axis.axhline(SOC_MIN, color="#7B8794", linestyle="--", linewidth=0.8)
    soc_axis.axhline(SOC_MAX, color="#7B8794", linestyle="--", linewidth=0.8)
    soc_axis.set_xlabel("自 2025-02-01 起的日序")
    soc_axis.set_ylabel("储电量 / kWh")
    soc_axis.legend(ncol=2)
    fig.tight_layout()
    save_high_res(fig, FIGURE_DIR / "problem2_cost_soc", formats=("pdf", "png"))


def strategy_summary(name, outputs):
    part = slice(OUTPUT_START, None)
    return {
        "name": name,
        "plan_purchase_kwh": float(outputs["grid"][part].sum()),
        "emergency_purchase_kwh": float(outputs["emergency"][part].sum()),
        "plan_cost_yuan": float(outputs["plan_cost"][part].sum()),
        "emergency_cost_yuan": float(outputs["emergency_cost"][part].sum()),
        "total_cost_yuan": float(outputs["total_cost"][part].sum()),
    }


def main():
    annual = load_problem2_data(DATA_PATH)
    typical = load_problem1_data(PRICE_PATH)
    dates = annual["dates"]
    price = np.asarray(typical["price"])
    load = annual["load_power"]
    pv = annual["pv_power"]

    load_models = generate_forecasts(load, dates, np.asarray(typical["load_power"]), seed=2025)
    pv_models = generate_forecasts(pv, dates, np.asarray(typical["pv_power"]), seed=2026)
    pv_models["ensemble"] = enforce_pv_physics(pv_models["ensemble"], pv)
    for name in pv_models["candidates"]:
        pv_models["candidates"][name] = enforce_pv_physics(pv_models["candidates"][name], pv)

    load_forecast = load_models["ensemble"]
    pv_forecast = pv_models["ensemble"]
    actual_net = load - pv
    forecast_net = load_forecast - pv_forecast
    margins = residual_margins(actual_net, forecast_net)

    metrics = {"load": {}, "pv": {}, "net": {}}
    for name in ["seasonal", "weekday", "ridge", "forest"]:
        metrics["load"][name] = prediction_metrics(load, load_models["candidates"][name])
        metrics["pv"][name] = prediction_metrics(pv, pv_models["candidates"][name])
        candidate_net = load_models["candidates"][name] - pv_models["candidates"][name]
        metrics["net"][name] = prediction_metrics(actual_net, candidate_net)
    metrics["load"]["ensemble"] = prediction_metrics(load, load_forecast)
    metrics["pv"]["ensemble"] = prediction_metrics(pv, pv_forecast)
    metrics["net"]["ensemble"] = prediction_metrics(actual_net, forecast_net)

    outputs = run_strategy(price, load, pv, load_forecast, pv_forecast, margins)
    checks_path = OUTPUT_DIR / "problem2_checks.json"
    if checks_path.exists():
        with checks_path.open(encoding="utf-8") as handle:
            comparisons = json.load(handle).get("strategy_comparison", [])
    else:
        comparisons = []
    if not comparisons:
        zero_margin = np.zeros_like(margins)
        seasonal_outputs = run_strategy(
            price,
            load,
            pv,
            load_models["candidates"]["seasonal"],
            pv_models["candidates"]["seasonal"],
            zero_margin,
        )
        point_outputs = run_strategy(price, load, pv, load_forecast, pv_forecast, zero_margin)
        comparisons = [
            strategy_summary("周同期点预测", seasonal_outputs),
            strategy_summary("机器学习集成点预测", point_outputs),
            strategy_summary("机器学习集成+80%分位裕度", outputs),
        ]
    else:
        comparisons[-1] = strategy_summary("机器学习集成+80%分位裕度", outputs)

    fill_result_workbook(dates, price, outputs)
    checks = write_outputs(dates, price, load, pv, load_forecast, pv_forecast, margins, outputs, metrics, comparisons)
    create_figures(dates, load, pv, load_forecast, pv_forecast, outputs, metrics)

    print(f"净负荷集成预测 MAE: {metrics['net']['ensemble']['mae']:.2f} kW")
    print(f"计划购电总量: {checks['total_plan_purchase_kwh']:.2f} kWh")
    print(f"紧急购电总量: {checks['total_emergency_purchase_kwh']:.2f} kWh")
    print(f"总购电费用: {checks['total_cost_yuan']:.2f} 元")
    print(f"储电量范围: [{checks['min_soc_kwh']:.2f}, {checks['max_soc_kwh']:.2f}] kWh")
    print(f"实际能量平衡最大残差: {checks['max_actual_balance_residual_kwh']:.3e} kWh")


if __name__ == "__main__":
    main()

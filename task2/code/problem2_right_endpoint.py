from __future__ import annotations

import csv
import json
import shutil
from copy import copy
from datetime import datetime
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook

import problem2 as base
import problem2_improved as improved
from plot_style import save_high_res


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_PATH = ROOT / "附件" / "附件5" / "result2.xlsx"
RESULT_PATH = ROOT / "result2_right_endpoint.xlsx"
OUTPUT_DIR = ROOT / "code" / "outputs"
FIGURE_DIR = ROOT / "figures"
LEFT_CHECKS_PATH = OUTPUT_DIR / "problem2_improved_checks.json"

N = 144


def to_right_endpoint(data):
    """Undo the former cyclic shift so 0:10 closes 0:00-0:10."""
    order = np.r_[np.arange(1, N), 0]
    result = dict(data)
    for key in ("price", "load", "pv", "load_forecast", "pv_forecast", "old_margin"):
        values = np.asarray(data[key])
        result[key] = values[order] if values.ndim == 1 else values[:, order]
    return result


def interval_label(slot):
    return f"{base.format_clock(slot * 10)}-{base.format_clock((slot + 1) * 10)}"


def apply_style(sheet, styles, row):
    for column, style in enumerate(styles, start=1):
        sheet.cell(row, column)._style = copy(style)


def write_workbook(data, outputs):
    shutil.copy2(TEMPLATE_PATH, RESULT_PATH)
    workbook = load_workbook(RESULT_PATH)
    plan = workbook["计划购电量"]
    for slot in range(N):
        plan.cell(1, slot + 2, interval_label(slot))
    for row, day in enumerate(range(len(data["dates"])), start=2):
        plan.cell(row, 1, datetime.combine(data["dates"][day], datetime.min.time()))
        for slot in range(N):
            plan.cell(row, slot + 2, float(outputs["grid"][day, slot]))
        plan.cell(row, 146, float(outputs["grid"][day].sum()))
        plan.cell(row, 147, float(outputs["plan_cost"][day]))

    storage = workbook["充放电量"]
    styles = [[copy(storage.cell(row, column)._style) for column in range(1, 7)] for row in range(2, 8)]
    storage.delete_rows(2, storage.max_row - 1)
    blocks = ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]
    row = 2
    for day, date in enumerate(data["dates"]):
        for block, label in enumerate(blocks):
            apply_style(storage, styles[block], row)
            start, stop = block * 24, (block + 1) * 24
            storage.cell(row, 1, datetime.combine(date, datetime.min.time()) if block == 0 else None)
            storage.cell(row, 2, label)
            storage.cell(row, 3, float(outputs["charge"][day, start:stop].sum()))
            storage.cell(row, 4, float(outputs["discharge"][day, start:stop].sum()))
            if block == 0:
                storage.cell(row, 5, "0:00")
                storage.cell(row, 6, float(outputs["soc"][day, 0]))
            elif block == 1:
                storage.cell(row, 5, "24:00")
                storage.cell(row, 6, float(outputs["soc"][day, -1]))
            row += 1

    emergency = workbook["紧急购电量"]
    style = [copy(emergency.cell(2, column)._style) for column in range(1, 4)]
    emergency.delete_rows(2, emergency.max_row - 1)
    row = 2
    for day, date in enumerate(data["dates"]):
        for index, (period, amount) in enumerate(base.emergency_segments(outputs["emergency"][day])):
            apply_style(emergency, style, row)
            emergency.cell(row, 1, datetime.combine(date, datetime.min.time()) if index == 0 else None)
            emergency.cell(row, 2, period)
            emergency.cell(row, 3, amount)
            row += 1
    workbook.save(RESULT_PATH)


def write_outputs(data, outputs):
    with (OUTPUT_DIR / "problem2_right_endpoint_schedule.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(["日期", "时段", "节点时刻", "电价", "实际负载电量", "实际光伏电量", "预测负载电量", "预测光伏电量", "自适应分位数", "安全裕度电量", "计划购电量", "实际充电量", "实际放电量", "紧急购电量", "弃电量", "时段末储电量"])
        for day, date in enumerate(data["dates"]):
            for slot in range(N):
                writer.writerow([date.isoformat(), interval_label(slot), base.format_clock((slot + 1) * 10), data["price"][slot], data["load"][day, slot], data["pv"][day, slot], data["load_forecast"][day, slot], data["pv_forecast"][day, slot], outputs["quantile"][day], outputs["margin"][day, slot], outputs["grid"][day, slot], outputs["charge"][day, slot], outputs["discharge"][day, slot], outputs["emergency"][day, slot], outputs["spill"][day, slot], outputs["soc"][day, slot + 1]])

    with (OUTPUT_DIR / "problem2_right_endpoint_daily_summary.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(["日期", "自适应分位数", "计划购电量(kWh)", "紧急购电量(kWh)", "计划购电费(元)", "紧急购电费(元)", "总费用(元)", "0:00储电量(kWh)", "24:00储电量(kWh)"])
        for day, date in enumerate(data["dates"]):
            writer.writerow([date.isoformat(), outputs["quantile"][day], outputs["grid"][day].sum(), outputs["emergency"][day].sum(), outputs["plan_cost"][day], outputs["emergency_cost"][day], outputs["total_cost"][day], outputs["soc"][day, 0], outputs["soc"][day, -1]])

    with LEFT_CHECKS_PATH.open(encoding="utf-8") as handle:
        left = json.load(handle)
    balance = outputs["grid"] + data["pv"] + outputs["discharge"] + outputs["emergency"] - data["load"] - outputs["charge"] - outputs["spill"]
    transition = outputs["soc"][:, 1:] - outputs["soc"][:, :-1] - base.ETA_CHARGE * outputs["charge"] + outputs["discharge"] / base.ETA_DISCHARGE
    right_total = float(outputs["total_cost"].sum())
    checks = {
        "time_interpretation": "right endpoint; 0:10 represents 0:00-0:10",
        "days": len(data["dates"]),
        "average_quantile": float(outputs["quantile"].mean()),
        "plan_purchase_kwh": float(outputs["grid"].sum()),
        "emergency_purchase_kwh": float(outputs["emergency"].sum()),
        "plan_cost_yuan": float(outputs["plan_cost"].sum()),
        "emergency_cost_yuan": float(outputs["emergency_cost"].sum()),
        "total_cost_yuan": right_total,
        "left_endpoint_total_cost_yuan": left["total_cost_yuan"],
        "right_minus_left_cost_yuan": right_total - left["total_cost_yuan"],
        "right_minus_left_cost_rate": right_total / left["total_cost_yuan"] - 1,
        "left_endpoint_emergency_purchase_kwh": left["emergency_purchase_kwh"],
        "right_minus_left_emergency_kwh": float(outputs["emergency"].sum()) - left["emergency_purchase_kwh"],
        "zero_emergency_days": int(np.sum(outputs["emergency"].sum(axis=1) <= 1e-6)),
        "emergency_intervals": int(np.sum(outputs["emergency"] > 1e-6)),
        "emergency_segments": int(sum(len(base.emergency_segments(day)) for day in outputs["emergency"])),
        "min_soc_kwh": float(outputs["soc"].min()),
        "max_soc_kwh": float(outputs["soc"].max()),
        "max_charge_kwh": float(outputs["charge"].max()),
        "max_discharge_kwh": float(outputs["discharge"].max()),
        "max_balance_residual_kwh": float(np.max(np.abs(balance))),
        "max_transition_residual_kwh": float(np.max(np.abs(transition))),
        "cross_day_soc_residual_kwh": float(np.max(np.abs(outputs["soc"][:-1, -1] - outputs["soc"][1:, 0]))),
    }

    workbook = load_workbook(RESULT_PATH, read_only=True, data_only=True)
    plan_rows = list(workbook["计划购电量"].iter_rows(min_row=2, values_only=True))
    workbook_plan = sum(sum(float(value or 0) for value in row[1:145]) for row in plan_rows)
    workbook_cost = sum(float(row[146] or 0) for row in plan_rows)
    emergency_rows = list(workbook["紧急购电量"].iter_rows(min_row=2, values_only=True))
    workbook_emergency = sum(float(row[2] or 0) for row in emergency_rows)
    checks.update({
        "workbook_plan_rows": len(plan_rows),
        "workbook_storage_rows": workbook["充放电量"].max_row - 1,
        "workbook_emergency_segments": len(emergency_rows),
        "workbook_plan_residual_kwh": workbook_plan - outputs["grid"].sum(),
        "workbook_cost_residual_yuan": workbook_cost - outputs["plan_cost"].sum(),
        "workbook_emergency_residual_kwh": workbook_emergency - outputs["emergency"].sum(),
    })
    with (OUTPUT_DIR / "problem2_right_endpoint_checks.json").open("w", encoding="utf-8") as handle:
        json.dump(checks, handle, ensure_ascii=False, indent=2)
    return checks


def create_figures(data, outputs):
    left_rows = list(csv.DictReader((OUTPUT_DIR / "problem2_improved_daily_summary.csv").open(encoding="utf-8-sig")))
    left_cost = np.asarray([float(row["总费用(元)"]) for row in left_rows])
    left_emergency = np.asarray([float(row["紧急购电量(kWh)"]) for row in left_rows])
    x = np.arange(len(data["dates"]))

    fig, (cost_axis, emergency_axis) = plt.subplots(2, 1, figsize=(11, 6.4), sharex=True)
    cost_axis.plot(x, left_cost, color="#7B8794", linewidth=0.9, label="左端点")
    cost_axis.plot(x, outputs["total_cost"], color="#2A7F8E", linewidth=0.9, label="右端点")
    cost_axis.set_ylabel("每日总费用 / 元")
    cost_axis.legend(ncol=2)
    emergency_axis.plot(x, left_emergency, color="#7B8794", linewidth=0.9, label="左端点")
    emergency_axis.plot(x, outputs["emergency"].sum(axis=1), color="#B44A3A", linewidth=0.9, label="右端点")
    emergency_axis.set_xlabel("自 2025-02-01 起的日序")
    emergency_axis.set_ylabel("紧急购电量 / kWh")
    emergency_axis.legend(ncol=2)
    fig.tight_layout()
    save_high_res(fig, FIGURE_DIR / "problem2_time_endpoint_sensitivity", formats=("pdf", "png"))


def main():
    data = to_right_endpoint(improved.load_cached_forecasts())
    outputs = improved.run_improved(data)
    write_workbook(data, outputs)
    checks = write_outputs(data, outputs)
    create_figures(data, outputs)
    print(f"右端点计划购电量: {checks['plan_purchase_kwh']:.2f} kWh")
    print(f"右端点紧急购电量: {checks['emergency_purchase_kwh']:.2f} kWh")
    print(f"右端点总费用: {checks['total_cost_yuan']:.2f} 元")
    print(f"相对左端点费用变化: {checks['right_minus_left_cost_yuan']:.2f} 元 ({checks['right_minus_left_cost_rate']:.4%})")


if __name__ == "__main__":
    main()

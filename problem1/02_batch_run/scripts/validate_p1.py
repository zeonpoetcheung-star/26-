#!/usr/bin/env python3
"""2026 C题 Problem1 A-5：独立复算 LP 结果、汇总及 result1 映射。"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import sys
from pathlib import Path
from typing import Any, Callable

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[4]
A_ROUTE = ROOT / "A_route"
COMMON = A_ROUTE / "common"
STAGE = A_ROUTE / "problem1" / "02_batch_run"
INPUT = ROOT / "input"
RESULTS = STAGE / "results"
TABLES = STAGE / "tables"
REPORTS = STAGE / "reports"
FIGURES = STAGE / "figures" / "diagnostic"
MODEL_PLAN = A_ROUTE / "problem1" / "01_model_plan" / "P1_MODEL_PLAN.md"
TASK_FILE = STAGE / "P1_BATCH_CODEX_TASK.md"

N = 144
DT_HOURS = 1.0 / 6.0
SOC_MIN = 1200.0
SOC_MAX = 10800.0
SOC_INITIAL = 6000.0
TRANSFER_MAX_KWH = 5000.0 / 6.0
ETA = 0.9
TOL = 1e-6
CHECK_COLUMNS = ["执行器", "检查编号", "阶段", "范围", "检查项", "状态", "阻塞性", "期望", "实际", "说明"]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def current_source_hashes() -> dict[str, str]:
    paths: list[Path] = [A_ROUTE / "CURRENT_STATE.md", MODEL_PLAN, TASK_FILE]
    for directory in (INPUT, COMMON):
        paths.extend(
            path for path in directory.rglob("*")
            if path.is_file() and "__pycache__" not in path.parts and path.suffix != ".pyc"
        )
    return {path.relative_to(ROOT).as_posix(): sha256(path) for path in sorted(set(paths))}


def close(a: Any, b: Any, tolerance: float = TOL) -> bool:
    return math.isclose(float(a), float(b), rel_tol=1e-10, abs_tol=tolerance)


def max_abs(values: np.ndarray) -> float:
    return float(np.max(np.abs(values))) if values.size else 0.0


class Checks:
    def __init__(self, phase: str) -> None:
        self.phase = phase
        self.rows: list[dict[str, Any]] = []
        self.index = 0

    def add(self, scope: str, item: str, expected: Any, actual: Any, passed: bool, *, blocking: bool = True, detail: str = "") -> None:
        self.index += 1
        self.rows.append({
            "执行器": "independent_validator",
            "检查编号": f"V{self.index:03d}",
            "阶段": self.phase,
            "范围": scope,
            "检查项": item,
            "状态": "PASS" if bool(passed) else "FAIL",
            "阻塞性": str(bool(blocking)).lower(),
            "期望": str(expected),
            "实际": str(actual),
            "说明": detail,
        })

    def guard(self, scope: str, item: str, fn: Callable[[], tuple[Any, Any, bool, str]], *, blocking: bool = True) -> None:
        try:
            expected, actual, passed, detail = fn()
            self.add(scope, item, expected, actual, passed, blocking=blocking, detail=detail)
        except Exception as exc:
            self.add(scope, item, "检查可完成", type(exc).__name__, False, blocking=blocking, detail=str(exc))


def expected_marker(slot_id: int) -> str:
    minutes = slot_id * 10
    if minutes == 1440:
        return "0:00+1"
    hours, minute = divmod(minutes, 60)
    return f"{hours}:{minute:02d}"


def expected_interval(slot_id: int) -> str:
    start = (slot_id - 1) * 10
    end = slot_id * 10
    start_text = f"{start // 60}:{start % 60:02d}"
    end_text = "24:00" if end == 1440 else f"{end // 60}:{end % 60:02d}"
    return f"{start_text}-{end_text}"


def read_workbook_values(path: Path) -> tuple[list[str], dict[str, list[list[Any]]], list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    names = workbook.sheetnames
    sheets: dict[str, list[list[Any]]] = {}
    error_cells: list[str] = []
    for name in names:
        sheet = workbook[name]
        rows = []
        for row in sheet.iter_rows():
            values = [cell.value for cell in row]
            rows.append(values)
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("#"):
                    error_cells.append(f"{name}!{cell.coordinate}:{cell.value}")
        sheets[name] = rows
    workbook.close()
    return names, sheets, error_cells


def validate_workbook(checks: Checks, schedule: pd.DataFrame, blocks: pd.DataFrame) -> None:
    output_path = RESULTS / "result1.xlsx"
    template_path = INPUT / "templates" / "result1.xlsx"
    checks.add("result1", "正式副本存在", True, output_path.exists(), output_path.is_file())
    if not output_path.is_file():
        return

    template_names, template_sheets, template_errors = read_workbook_values(template_path)
    output_names, output_sheets, output_errors = read_workbook_values(output_path)
    checks.add("result1", "工作表名称与顺序", template_names, output_names, output_names == template_names == ["计划购电量", "充放电量"])
    checks.add("result1", "无公式错误值", "0 个", len(output_errors), not output_errors, detail=", ".join(output_errors))

    plan_template = template_sheets["计划购电量"]
    plan_output = output_sheets["计划购电量"]
    storage_template = template_sheets["充放电量"]
    storage_output = output_sheets["充放电量"]
    labels_same = [row[0] for row in plan_output[:145]] == [row[0] for row in plan_template[:145]]
    checks.add("result1", "计划购电量标签未改", "144/144", f"{sum(a[0] == b[0] for a, b in zip(plan_template[:145], plan_output[:145]))}/145 含表头", labels_same)

    output_purchase = np.array([plan_output[row][1] for row in range(1, 145)], dtype=float)
    internal_purchase = schedule.sort_values("slot_id")["purchase_kwh"].to_numpy(dtype=float)
    expected_purchase = np.r_[internal_purchase[1:], internal_purchase[0]]
    mapping_error = max_abs(output_purchase - expected_purchase)
    checks.add("result1", "计划购电量循环映射", "最大误差<1e-6", mapping_error, mapping_error < TOL)
    checks.add("result1", "模板首行取内部第二槽", float(internal_purchase[1]), float(output_purchase[0]), close(output_purchase[0], internal_purchase[1]))
    checks.add("result1", "模板末行取下一重复日第一槽", float(internal_purchase[0]), float(output_purchase[-1]), close(output_purchase[-1], internal_purchase[0]))

    intended_plan = {(row, 1) for row in range(1, 145)}
    unchanged_plan = True
    for row in range(max(len(plan_template), len(plan_output))):
        width = max(len(plan_template[row]) if row < len(plan_template) else 0, len(plan_output[row]) if row < len(plan_output) else 0)
        for col in range(width):
            if (row, col) in intended_plan:
                continue
            old = plan_template[row][col] if row < len(plan_template) and col < len(plan_template[row]) else None
            new = plan_output[row][col] if row < len(plan_output) and col < len(plan_output[row]) else None
            unchanged_plan = unchanged_plan and old == new

    intended_storage = {(row, col) for row in range(1, 7) for col in (1, 2)} | {(1, 4), (2, 4)}
    unchanged_storage = True
    for row in range(max(len(storage_template), len(storage_output))):
        width = max(len(storage_template[row]) if row < len(storage_template) else 0, len(storage_output[row]) if row < len(storage_output) else 0)
        for col in range(width):
            if (row, col) in intended_storage:
                continue
            old = storage_template[row][col] if row < len(storage_template) and col < len(storage_template[row]) else None
            new = storage_output[row][col] if row < len(storage_output) and col < len(storage_output[row]) else None
            unchanged_storage = unchanged_storage and old == new
    checks.add("result1", "仅填写目标数值单元格", True, unchanged_plan and unchanged_storage, unchanged_plan and unchanged_storage)

    output_charge = np.array([storage_output[row][1] for row in range(1, 7)], dtype=float)
    output_discharge = np.array([storage_output[row][2] for row in range(1, 7)], dtype=float)
    charge_error = max_abs(output_charge - blocks["充电量_kWh"].to_numpy(dtype=float))
    discharge_error = max_abs(output_discharge - blocks["放电量_kWh"].to_numpy(dtype=float))
    checks.add("result1", "六个四小时充放电量映射", "最大误差<1e-6", max(charge_error, discharge_error), max(charge_error, discharge_error) < TOL)
    state_values = (storage_output[1][4], storage_output[2][4])
    checks.add("result1", "0:00与24:00储电量", "6000,6000", state_values, close(state_values[0], 6000) and close(state_values[1], 6000))
    checks.add("result1", "官方模板本体未修改", sha256(template_path), checks_source_hash(template_path), sha256(template_path) == checks_source_hash(template_path))


_EXPECTED_SOURCE_HASHES: dict[str, str] = {}


def checks_source_hash(path: Path) -> str:
    relative = path.relative_to(ROOT).as_posix()
    return _EXPECTED_SOURCE_HASHES.get(relative, "未记录")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--phase", choices=["pre_export", "final"], required=True)
    args = parser.parse_args()
    phase_cn = "导出前" if args.phase == "pre_export" else "最终"
    checks = Checks(phase_cn)

    required = [
        RESULTS / "p1_schedule_internal.csv",
        RESULTS / "p1_summary.json",
        RESULTS / "p1_efficiency_sensitivity.csv",
        TABLES / "p1_specified_purchase_intervals.csv",
        TABLES / "p1_storage_4h_summary.csv",
        TABLES / "p1_baseline_comparison.csv",
    ]
    missing = [path.relative_to(STAGE).as_posix() for path in required if not path.is_file()]
    checks.add("产物", "canonical 产物齐全", "缺失 0", len(missing), not missing, detail=", ".join(missing))
    if missing:
        TABLES.mkdir(parents=True, exist_ok=True)
        pd.DataFrame(checks.rows, columns=CHECK_COLUMNS).to_csv(TABLES / "p1_validation_checks.csv", index=False, encoding="utf-8-sig", lineterminator="\n")
        return 1

    summary = json.loads((RESULTS / "p1_summary.json").read_text(encoding="utf-8"))
    global _EXPECTED_SOURCE_HASHES
    _EXPECTED_SOURCE_HASHES = dict(summary["source_hashes"])
    current_hashes = current_source_hashes()
    missing_sources = sorted(set(_EXPECTED_SOURCE_HASHES) - set(current_hashes))
    added_sources = sorted(set(current_hashes) - set(_EXPECTED_SOURCE_HASHES))
    changed_sources = sorted(key for key in set(current_hashes) & set(_EXPECTED_SOURCE_HASHES) if current_hashes[key] != _EXPECTED_SOURCE_HASHES[key])
    hash_ok = not missing_sources and not added_sources and not changed_sources
    checks.add("冻结源", "Common/input/模型计划未变", "缺失/新增/变更均为0", f"缺失={len(missing_sources)},新增={len(added_sources)},变更={len(changed_sources)}", hash_ok, detail=str({"缺失": missing_sources, "新增": added_sources, "变更": changed_sources}))

    gate_exists = (REPORTS / "P1_GATE.md").exists()
    checks.add("流水线", "校验时阶段验收文件尚未生成", False, gate_exists, not gate_exists)
    result_exists = (RESULTS / "result1.xlsx").exists()
    if args.phase == "pre_export":
        checks.add("流水线", "导出前正式副本尚未写入", False, result_exists, not result_exists)

    forbidden_dirs = [STAGE / name for name in ("model", "models", "optimization")]
    present_forbidden = [path.name for path in forbidden_dirs if path.exists()]
    checks.add("范围", "未创建模型搜索或优化目录", "无", present_forbidden, not present_forbidden)

    schedule = pd.read_csv(RESULTS / "p1_schedule_internal.csv", encoding="utf-8-sig")
    blocks = pd.read_csv(TABLES / "p1_storage_4h_summary.csv", encoding="utf-8-sig")
    specified = pd.read_csv(TABLES / "p1_specified_purchase_intervals.csv", encoding="utf-8-sig")
    baseline = pd.read_csv(TABLES / "p1_baseline_comparison.csv", encoding="utf-8-sig")
    sensitivity = pd.read_csv(RESULTS / "p1_efficiency_sensitivity.csv", encoding="utf-8-sig")

    required_columns = {
        "slot_id", "source_marker", "physical_interval", "duration_hours", "price_yuan_per_kwh",
        "load_kw", "pv_forecast_kw", "load_kwh", "pv_forecast_kwh", "purchase_kwh",
        "charge_bus_kwh", "discharge_bus_kwh", "curtailment_kwh", "soc_start_kwh",
        "soc_end_kwh", "purchase_cost_yuan",
    }
    columns_ok = required_columns.issubset(schedule.columns)
    checks.add("结构", "内部 schedule 字段", "所需字段齐全", sorted(required_columns - set(schedule.columns)), columns_ok)
    structure_ok = len(schedule) == N and schedule["slot_id"].astype(int).tolist() == list(range(1, N + 1)) and not schedule["slot_id"].duplicated().any()
    checks.add("结构", "144 个顺序槽位", "1..144且唯一", len(schedule), structure_ok)

    marker_ok = schedule["source_marker"].astype(str).tolist() == [expected_marker(slot) for slot in range(1, N + 1)]
    interval_ok = schedule["physical_interval"].astype(str).tolist() == [expected_interval(slot) for slot in range(1, N + 1)]
    duration_error = max_abs(schedule["duration_hours"].to_numpy(dtype=float) - DT_HOURS)
    checks.add("时间", "H-END source marker", "144/144", int(marker_ok) * 144, marker_ok)
    checks.add("时间", "物理区间完整覆盖0:00-24:00", "144/144连续", int(interval_ok) * 144, interval_ok)
    checks.add("时间", "每槽时长", "1/6 h", duration_error, duration_error < TOL)

    load_conversion_error = max_abs(schedule["load_kwh"].to_numpy(dtype=float) - schedule["load_kw"].to_numpy(dtype=float) / 6.0)
    pv_conversion_error = max_abs(schedule["pv_forecast_kwh"].to_numpy(dtype=float) - schedule["pv_forecast_kw"].to_numpy(dtype=float) / 6.0)
    checks.add("单位", "kW转kWh乘1/6", "最大误差<1e-6", max(load_conversion_error, pv_conversion_error), max(load_conversion_error, pv_conversion_error) < TOL)

    g = schedule["purchase_kwh"].to_numpy(dtype=float)
    c = schedule["charge_bus_kwh"].to_numpy(dtype=float)
    d = schedule["discharge_bus_kwh"].to_numpy(dtype=float)
    w = schedule["curtailment_kwh"].to_numpy(dtype=float)
    load = schedule["load_kwh"].to_numpy(dtype=float)
    pv = schedule["pv_forecast_kwh"].to_numpy(dtype=float)
    price = schedule["price_yuan_per_kwh"].to_numpy(dtype=float)
    soc_start = schedule["soc_start_kwh"].to_numpy(dtype=float)
    soc_end = schedule["soc_end_kwh"].to_numpy(dtype=float)
    soc = np.r_[soc_start[0], soc_end]

    balance_residual = g + pv + d - load - c - w
    soc_residual = soc_end - soc_start - ETA * c + d / ETA
    continuity_error = max_abs(soc_start[1:] - soc_end[:-1])
    checks.add("求解器", "连续LP最优", "optimal", summary.get("solver_status"), summary.get("solver_success") is True and summary.get("solver_status") == "optimal")
    checks.add("能量", "144槽能量守恒", "最大绝对误差<1e-6 kWh", max_abs(balance_residual), max_abs(balance_residual) < TOL)
    checks.add("SOC", "145个状态点且连续", "145且连续误差<1e-6", f"145, {continuity_error}", len(soc) == 145 and continuity_error < TOL)
    checks.add("SOC", "144槽状态递推", "最大绝对误差<1e-6 kWh", max_abs(soc_residual), max_abs(soc_residual) < TOL)
    soc_range_ok = soc.min() >= SOC_MIN - TOL and soc.max() <= SOC_MAX + TOL
    checks.add("SOC", "状态范围", "1200..10800 kWh", f"{soc.min()}..{soc.max()}", soc_range_ok)
    checks.add("SOC", "S0=S144=6000", "6000,6000", f"{soc[0]},{soc[-1]}", close(soc[0], SOC_INITIAL) and close(soc[-1], SOC_INITIAL))

    power_ok = c.min() >= -TOL and d.min() >= -TOL and c.max() <= TRANSFER_MAX_KWH + TOL and d.max() <= TRANSFER_MAX_KWH + TOL
    checks.add("储能", "充放电量限制", f"0..{TRANSFER_MAX_KWH}", f"C={c.min()}..{c.max()},D={d.min()}..{d.max()}", power_ok)
    simultaneous = int(np.sum((c > TOL) & (d > TOL)))
    checks.add("储能", "无实质同时充放电", 0, simultaneous, simultaneous == 0)
    grid_pv_ok = g.min() >= -TOL and w.min() >= -TOL and np.all(w <= pv + TOL)
    checks.add("外网与光伏", "购电非负且弃光在范围内", "G>=0,0<=W<=PV", f"minG={g.min()},minW={w.min()},max(W-PV)={(w-pv).max()}", grid_pv_ok)
    checks.add("模型边界", "决策变量集合无售电变量", ["G", "C", "D", "W", "S"], summary.get("model_variables"), summary.get("model_variables") == ["G", "C", "D", "W", "S"])

    cost_by_rows = price * g
    row_cost_error = max_abs(cost_by_rows - schedule["purchase_cost_yuan"].to_numpy(dtype=float))
    cost_total = float(cost_by_rows.sum())
    checks.add("成本", "单槽成本独立复算", "最大误差<1e-6", row_cost_error, row_cost_error < TOL)
    checks.add("成本", "全天成本独立复算", summary["total_purchase_cost_yuan"], cost_total, close(cost_total, summary["total_purchase_cost_yuan"]))

    baseline_g = np.maximum(load - pv, 0.0)
    baseline_w = np.maximum(pv - load, 0.0)
    baseline_cost = float(price @ baseline_g)
    baseline_row = baseline[baseline["方案"] == "无储能基准"].iloc[0]
    baseline_ok = close(baseline_row["全天购电量_kWh"], baseline_g.sum()) and close(baseline_row["全天购电费_元"], baseline_cost) and close(baseline_row["全天弃光量_kWh"], baseline_w.sum())
    checks.add("基准", "无储能基准独立复算", "购电量/成本/弃光一致", baseline_ok, baseline_ok)
    checks.add("基准", "LP成本不高于基准", f"<={baseline_cost + TOL}", cost_total, cost_total <= baseline_cost + TOL)

    totals_ok = close(g.sum(), summary["total_purchase_kwh"]) and close(cost_total, summary["total_purchase_cost_yuan"]) and close(c.sum(), summary["total_charge_bus_kwh"]) and close(d.sum(), summary["total_discharge_bus_kwh"])
    checks.add("汇总", "144槽全天总量", "购电/成本/充电/放电一致", totals_ok, totals_ok)
    expected_block_labels = ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]
    block_ok = len(blocks) == 6 and blocks["时间段"].astype(str).tolist() == expected_block_labels
    block_error = 0.0
    if block_ok:
        expected_c = np.array([c[i*24:(i+1)*24].sum() for i in range(6)])
        expected_d = np.array([d[i*24:(i+1)*24].sum() for i in range(6)])
        block_error = max(max_abs(blocks["充电量_kWh"].to_numpy(dtype=float) - expected_c), max_abs(blocks["放电量_kWh"].to_numpy(dtype=float) - expected_d))
        block_ok = block_error < TOL and close(blocks["充电量_kWh"].sum(), c.sum()) and close(blocks["放电量_kWh"].sum(), d.sum())
    checks.add("汇总", "六个四小时块", "分块与全天汇总一致", block_error, block_ok)

    specified_slots = [61, 73, 85, 97, 109, 121]
    specified_ok = len(specified) == 6 and specified["slot_id"].astype(int).tolist() == specified_slots
    specified_error = float("inf")
    if specified_ok:
        specified_error = max_abs(specified["计划购电量_kWh"].to_numpy(dtype=float) - g[np.array(specified_slots) - 1])
        specified_ok = specified_error < TOL
    checks.add("汇总", "六个指定购电时段H-END映射", specified_slots, f"最大误差={specified_error}", specified_ok)

    sensitivity_ids = sensitivity["scenario_id"].astype(str).tolist()
    sensitivity_ok = len(sensitivity) == 2 and sensitivity_ids == ["main_eta_0p9", "round_trip_eta_sqrt_0p9"]
    if sensitivity_ok:
        main_row, alt_row = sensitivity.iloc[0], sensitivity.iloc[1]
        sensitivity_ok = (
            close(main_row["eta_c"], 0.9)
            and close(main_row["eta_d"], 0.9)
            and close(alt_row["eta_c"], math.sqrt(0.9))
            and close(alt_row["eta_d"], math.sqrt(0.9))
            and main_row["solver_status"] == "optimal"
            and alt_row["solver_status"] == "optimal"
            and close(main_row["全天购电费_元"], cost_total)
        )
    checks.add("敏感性", "仅主口径与sqrt(0.9)口径", "2个固定口径且均optimal", sensitivity_ids, sensitivity_ok)

    figure_names = sorted(path.name for path in FIGURES.glob("*.png"))
    expected_figures = sorted(["p1_purchase_and_price.png", "p1_soc_trajectory.png", "p1_charge_discharge.png"])
    checks.add("图形", "内部诊断图数量和名称", expected_figures, figure_names, figure_names == expected_figures)

    if args.phase == "final":
        validate_workbook(checks, schedule, blocks)

    check_frame = pd.DataFrame(checks.rows, columns=CHECK_COLUMNS)
    TABLES.mkdir(parents=True, exist_ok=True)
    check_frame.to_csv(TABLES / "p1_validation_checks.csv", index=False, encoding="utf-8-sig", lineterminator="\n")
    pass_count = int((check_frame["状态"] == "PASS").sum())
    fail_frame = check_frame[check_frame["状态"] == "FAIL"]
    blocking_fail = fail_frame[fail_frame["阻塞性"].astype(str).str.lower() == "true"]

    if args.phase == "final":
        REPORTS.mkdir(parents=True, exist_ok=True)
        REPORTS.joinpath("P1_VALIDATION_REPORT.md").write_text(f"""# Problem1 A-5 独立校验报告

## 结论

- 校验范围：canonical 结果、时间与单位、能量守恒、SOC、约束、成本、基准、汇总、敏感性及 `result1.xlsx`。
- PASS：{pass_count} 项。
- FAIL：{len(fail_frame)} 项。
- 阻塞失败：{len(blocking_fail)} 项。

## 失败项

{chr(10).join(f"- {row['检查编号']}：{row['检查项']}（期望：{row['期望']}；实际：{row['实际']}；说明：{row['说明']}）" for _, row in fail_frame.iterrows()) if not fail_frame.empty else '- 无。'}

全部数值检查由本脚本从 canonical 输入和导出文件独立复算，没有调用求解脚本中的计算函数。
""", encoding="utf-8", newline="\n")

    print(f"{phase_cn}独立校验：PASS={pass_count}，FAIL={len(fail_frame)}，阻塞失败={len(blocking_fail)}")
    if not fail_frame.empty:
        for _, row in fail_frame.iterrows():
            print(f"{row['检查编号']} {row['范围']} / {row['检查项']}：{row['说明']}")
    return 1 if not blocking_fail.empty else 0


if __name__ == "__main__":
    raise SystemExit(main())

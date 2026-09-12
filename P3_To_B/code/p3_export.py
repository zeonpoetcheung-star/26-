"""独立数值验证后首次导出MAIN，随后openpyxl只读核对原模板和全部输出单元格。"""
import os,subprocess,datetime
from pathlib import Path
import numpy as np
import pandas as pd
import openpyxl
from p3_io import save_json,read_json,file_hash,write_csv

def excel_date(value):
    return (datetime.datetime.fromisoformat(value)-datetime.datetime(1899,12,30)).days

def clock(slot):
    minute=int(slot)*10
    return f'{minute//60}:{minute%60:02d}'

def export_and_validate(root,base,spec):
    destination=base/'results/result3.xlsx';evidence=base/'logs/workbook_validation.json'
    audited=read_json(base/'logs/independent_validation.json')
    causal=read_json(base/'logs/causality_validation.json')
    if audited['blocking'] or not audited['annual_complete'] or causal['FAIL']:
        raise RuntimeError('独立数值/因果验证未全部通过，禁止首次导出')
    if destination.exists():
        if not evidence.exists(): raise RuntimeError('工作簿已存在但缺独立验收；不得静默重导')
        old=read_json(evidence)
        if file_hash(destination)!=old['sha256']: raise RuntimeError('已导出工作簿身份改变')
        return old
    main=spec['primary_policy'];actual=pd.read_csv(base/'results/policies'/main/'p3_actual_schedule.csv')
    daily=pd.read_csv(base/'results/p3_daily_summary.csv');daily=daily[daily.policy_id==main].sort_values('date')
    storage=pd.read_csv(base/'tables/p3_storage_4h_summary.csv')
    emergency=pd.read_csv(base/'results/p3_emergency_events.csv');emergency=emergency[emergency.policy_id==main].sort_values(['date','start_slot'])
    payload={'pre_export_validated':True,'policy_id':main,'plan':[],'adjust':[],'storage':[],'emergency':[]}
    for date,frame in actual.groupby('date',sort=True):
        payload['plan'].append(frame.g0_kwh.tolist()+[float(frame.g0_kwh.sum()),float(np.dot(frame.g0_kwh,frame.price_yuan_per_kwh))])
        payload['adjust'].append(frame.a_kwh.tolist()+[float(frame.a_kwh.sum()),float(frame.total_cost_yuan.sum())])
        ds=daily[daily.date==date].iloc[0]
        for b,row in enumerate(storage[storage.date==date].itertuples()):
            payload['storage'].append([excel_date(date) if b==0 else None,row.time_interval,row.charge_bus_kwh,row.discharge_bus_kwh,
                 0 if b==0 else ('24:00' if b==1 else None),float(ds.soc_start_kwh) if b==0 else (float(ds.soc_end_kwh) if b==1 else None)])
    for row in emergency.itertuples():
        payload['emergency'].append([excel_date(row.date),clock(row.start_slot-1)+'-'+clock(row.end_slot),row.emergency_kwh])
    save_json(base/'logs/export_payload.json',payload)
    node=os.environ.get('P3_NODE_EXE');marker=os.environ.get('P3_ARTIFACT_MARKER')
    if not node or not marker: raise RuntimeError('尚未提供经workspace依赖定位的Node/工作簿操作标记路径')
    subprocess.run([node,marker,'--operation-kind','create','--expected-output-count','1','--output-format','xlsx'],check=True,cwd=base)
    subprocess.run([node,str(base/'scripts/p3_export_workbook.mjs'),str(root),'export'],check=True,cwd=base)
    template=openpyxl.load_workbook(root/'input/templates/result3.xlsx',read_only=False,data_only=False)
    book=openpyxl.load_workbook(destination,read_only=False,data_only=False)
    checks=[];mapping=[]
    def check(name,ok,detail='',residual=0.): checks.append({'check_id':name,'status':'PASS' if bool(ok) else 'FAIL','detail':detail,'max_residual':float(residual)})
    check('XLSX_SHEETS_PRESERVED',book.sheetnames==template.sheetnames)
    for s in range(4):
        expected=[c.value for c in template.worksheets[s][1]];got=[book.worksheets[s].cell(1,j+1).value for j in range(len(expected))]
        check('XLSX_HEADERS_'+str(s),got==expected,'保留全部原始标签；H-END物理映射另列')
    for s,name in enumerate(['plan','adjust']):
        sheet=book.worksheets[s];matrix=payload[name];worst=0.
        check('XLSX_DAYS_'+name,sheet.max_row==335 and sheet.max_column==147)
        for d,row in enumerate(matrix):
            date=daily.date.iloc[d]
            check('XLSX_DATE_'+name+'_'+date,str(sheet.cell(d+2,1).value)[:10]==date)
            for j,val in enumerate(row):
                cell=sheet.cell(d+2,j+2);numeric=isinstance(cell.value,(float,int)) and not isinstance(cell.value,bool)
                residual=abs(cell.value-val) if numeric else float('inf');worst=max(worst,residual)
                mapping.append({'sheet':sheet.title,'cell':cell.coordinate,'date':date,'slot_id':j+1 if j<144 else '',
                    'meaning':('0时g0' if s==0 else '交付前最终a') if j<144 else ('总普通电量' if j==144 else ('初始报价' if s==0 else '含紧急的真实总账')),
                    'source_file':'p3_initial_plan.csv' if s==0 else 'p3_settlement_ledger.csv','expected_value':val,'observed_value':cell.value,'max_residual':residual})
        check('XLSX_ALL_VALUES_'+name,worst<=1e-6,f'{len(matrix)*144}个购电量和668个日汇总单元格',worst)
    st=book.worksheets[2];worst=0.
    check('XLSX_STORAGE_BLOCKS',st.max_row==2005,'2004个四小时块及334对SOC')
    for i,row in enumerate(payload['storage']):
        for j in (2,3,5):
            if row[j] is None: continue
            value=st.cell(i+2,j+1).value;res=abs(float(value)-row[j]) if isinstance(value,(int,float)) else float('inf');worst=max(worst,res)
            mapping.append({'sheet':st.title,'cell':st.cell(i+2,j+1).coordinate,'date':storage.date.iloc[i],'slot_id':'',
                            'meaning':['日期','时段','真实充电','真实放电','时刻','真实边界SOC'][j],
                            'source_file':'p3_storage_4h_summary.csv','expected_value':row[j],'observed_value':value,'max_residual':res})
        if i%6==0: check('XLSX_STORAGE_DATE_'+str(i),str(st.cell(i+2,1).value)[:10]==storage.date.iloc[i])
        check('XLSX_STORAGE_INTERVAL_'+str(i),st.cell(i+2,2).value==row[1])
    check('XLSX_STORAGE_VALUES',worst<=1e-6,residual=worst)
    es=book.worksheets[3];worst=0.
    check('XLSX_EVENT_ROWS',es.max_row==max(11,len(payload['emergency'])+1),'模板已有空白尾部允许保留；事件不含阈值以下残余')
    for i,row in enumerate(payload['emergency']):
        orig=emergency.iloc[i];value=es.cell(i+2,3).value;res=abs(float(value)-row[2]) if isinstance(value,(int,float)) else float('inf');worst=max(worst,res)
        check('XLSX_EMERGENCY_KEY_'+str(i),str(es.cell(i+2,1).value)[:10]==orig.date and es.cell(i+2,2).value==row[1])
        mapping.append({'sheet':es.title,'cell':es.cell(i+2,3).coordinate,'date':orig.date,'slot_id':f'{orig.start_slot}-{orig.end_slot}',
                        'meaning':'真实连续紧急事件电量','source_file':'p3_emergency_events.csv','expected_value':row[2],'observed_value':value,'max_residual':res})
    check('XLSX_EMERGENCY_VALUES',worst<=1e-6,residual=worst)
    formula_errors=[]
    for sheet in book:
        for row in sheet:
            for cell in row:
                if cell.data_type=='e': formula_errors.append(f'{sheet.title}!{cell.coordinate}:{cell.value}')
    check('XLSX_NO_CELL_ERRORS',not formula_errors,str(formula_errors[:10]))
    template.close();book.close()
    write_csv(base/'tables/p3_output_mapping_audit.csv',mapping)
    result={'sha256':file_hash(destination),'PASS':sum(r['status']=='PASS' for r in checks),'FAIL':sum(r['status']=='FAIL' for r in checks),
            'checks':checks,'workbook_policy':main,'export_count':1}
    save_json(evidence,result)
    if result['FAIL']: raise RuntimeError('工作簿独立读回失败；保留首次导出，不静默覆盖')
    return result
